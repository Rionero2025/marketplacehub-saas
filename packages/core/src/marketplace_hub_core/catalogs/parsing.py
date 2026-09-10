from __future__ import annotations

import csv
import io
import json
import math
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePath
from typing import Any

import xlrd
from defusedxml import ElementTree as SafeElementTree
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook

from marketplace_hub_core.catalogs.schema import MAX_CATALOG_ARTIFACT_BYTES

MAX_PRODUCTS = 200_000
MAX_COLUMNS = 300
MAX_CELLS = 5_000_000
MAX_CELL_CHARS = 20_000
MAX_XLSX_ENTRIES = 2_000
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_XML_ELEMENTS = 1_000_000


class CatalogFileLimitError(ValueError):
    pass


class CatalogFileValidationError(ValueError):
    pass


ALIASES = {
    "ean": (
        "ean", "ean13", "ean_13", "barcode", "bar_code", "gtin", "codice_ean",
        "kod_ean", "item_ean", "itemean",
    ),
    "sku": (
        "sku", "tag_sku", "seller_sku", "reference", "referencia", "ref", "codice",
        "product_id", "ref_proveedor", "product_code", "item_code", "item_part_number",
        "itempartnumber", "erp_id", "erpid", "symbol", "index", "indeks", "kod",
        "kod_towaru", "catalog_number", "catalogue_number",
    ),
    "name": (
        "name", "title", "nome", "prodotto", "product_name", "description", "titulo",
        "nazwa", "nazwa_towaru", "nazwa_produktu", "nombre_completo",
        "nombre_producto___modelo__it_", "nombre_producto___modelo",
    ),
    "cost": (
        "neto_italia_(zona_2)", "neto_italia", "neto_pt", "cost", "costo", "net_price",
        "wholesale_price", "price_net", "price_nett", "pricenett", "pricelist_eur",
        "purchase_price", "pvd", "cena_netto", "cena_zakupu", "cena_hurtowa",
        "your_price", "customer_price", "prezzo_acquisto", "prezzo_di_acquisto",
    ),
    "shipping_cost": (
        "shipping_cost", "delivery_cost", "freight_cost", "costo_spedizione",
        "costo_di_spedizione", "spedizione",
    ),
    "total_cost": (
        "total_cost", "landed_cost", "costo_totale", "costo_totale_acquisto",
    ),
    "quantity": (
        "quantity", "qty", "stock_prod", "stock_mwg", "stock", "availability",
        "available_qty", "availableqty", "available_quantity", "qta", "stan",
        "stan_magazynowy", "ilosc", "dostepnosc", "disponibilita", "quantita",
    ),
}


def _token(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", "_", text.casefold()).strip("_")


ALIAS_TOKENS = {
    target: {_token(value) for value in (target, *values)} for target, values in ALIASES.items()
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
    result = str(value).strip()
    return "" if result.casefold() in {"nan", "none", "null"} else result


def _excel_text(value: Any, number_format: str = "") -> str:
    if isinstance(value, int | float) and not isinstance(value, bool):
        pattern = str(number_format or "").split(";", 1)[0].strip()
        if re.fullmatch(r"0+", pattern) and float(value).is_integer():
            return str(int(value)).zfill(len(pattern))
    return _text(value)


def _decimal(value: Any) -> Decimal | None:
    text = _text(value).replace("\u00a0", "").replace(" ", "")
    if not text:
        return None
    text = re.sub(r"[^0-9,\.\-+]", "", text)
    if not text or text in {"+", "-", ".", ","}:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        result = Decimal(text)
    except InvalidOperation:
        return None
    if not result.is_finite() or abs(result) >= Decimal("1e30"):
        return None
    return result


def _safe_json_value(value: Any) -> str:
    result = _text(value)
    if len(result) > MAX_CELL_CHARS:
        raise CatalogFileLimitError("Il listino contiene un valore troppo lungo.")
    return result


def _header_map(headers: list[Any]) -> dict[str, int]:
    if len(headers) > MAX_COLUMNS:
        raise CatalogFileLimitError("Il listino contiene troppe colonne.")
    tokens = [_token(value) for value in headers]
    result: dict[str, int] = {}
    for target, aliases in ALIAS_TOKENS.items():
        for index, token in enumerate(tokens):
            if token in aliases:
                result[target] = index
                break
    if not result:
        raise CatalogFileValidationError(
            "Il listino non contiene colonne riconoscibili per prodotto, codice o costo."
        )
    return result


def _normalized_rows(
    headers: list[Any],
    values: list[list[Any]],
    *,
    source_row_start: int,
) -> list[dict[str, Any]]:
    mapping = _header_map(headers)
    # Some Hurtel exports put a supplier reference under EAN and the real GTIN
    # under SKU. Preserve the legacy whole-feed repair, but only when the shape
    # evidence is strong enough to avoid swapping isolated malformed values.
    def ean_ratio(index: int | None) -> float:
        if index is None:
            return 0.0
        populated = [
            _text(row[index]) for row in values if index < len(row) and _text(row[index])
        ]
        if not populated:
            return 0.0
        matching = sum(
            bool(re.fullmatch(r"(?:\d{8}|\d{12,14})(?:\.0+)?", item))
            for item in populated
        )
        return matching / len(populated)

    ean_index = mapping.get("ean")
    sku_index = mapping.get("sku")
    if ean_index is not None and sku_index is not None:
        if ean_ratio(ean_index) < 0.25 and ean_ratio(sku_index) > 0.75:
            mapping["ean"], mapping["sku"] = sku_index, ean_index
    header_names = []
    used: dict[str, int] = {}
    for index, header in enumerate(headers):
        name = _token(header) or f"colonna_{index + 1}"
        used[name] = used.get(name, 0) + 1
        header_names.append(name if used[name] == 1 else f"{name}_{used[name]}")

    output: list[dict[str, Any]] = []
    for source_row, row in enumerate(values, start=source_row_start):
        if len(output) >= MAX_PRODUCTS:
            raise CatalogFileLimitError("Il listino contiene troppi prodotti.")
        cells = list(row[:MAX_COLUMNS])
        if not any(_text(value) for value in cells):
            continue
        source = {
            header_names[index]: _safe_json_value(value)
            for index, value in enumerate(cells)
            if _safe_json_value(value)
        }

        def value(target: str) -> Any:
            index = mapping.get(target)
            return cells[index] if index is not None and index < len(cells) else None

        ean = _text(value("ean"))
        if re.fullmatch(r"\d+\.0+", ean):
            ean = ean.split(".", 1)[0]
        sku = _text(value("sku"))
        name = _text(value("name"))
        cost = _decimal(value("cost"))
        shipping = _decimal(value("shipping_cost"))
        total = _decimal(value("total_cost"))
        quantity = _decimal(value("quantity"))
        cost = cost if cost is not None else Decimal("0")
        shipping = shipping if shipping is not None else Decimal("0")
        total = total if total is not None else cost + shipping
        quantity = quantity if quantity is not None else Decimal("0")
        canonical = {
            "ean": ean,
            "sku": sku,
            "name": name,
            "cost": format(cost, "f"),
            "shipping_cost": format(shipping, "f"),
            "total_cost": format(total, "f"),
            "quantity": format(quantity, "f"),
            "source": source,
        }
        output.append({
            "source_row": source_row,
            "ean": ean,
            "sku": sku,
            "name": name,
            "cost": cost,
            "shipping_cost": shipping,
            "total_cost": total,
            "quantity": quantity,
            "canonical_json": json.dumps(
                canonical, ensure_ascii=False, allow_nan=False, separators=(",", ":"),
            ),
        })
    if not output:
        raise CatalogFileValidationError("Il listino non contiene prodotti leggibili.")
    return output


def _decode_text(content: bytes) -> str:
    if b"\x00" in content:
        raise CatalogFileValidationError("Il file di testo non è valido.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise CatalogFileValidationError("La codifica del file di testo non è supportata.")


def _parse_delimited(content: bytes, file_format: str) -> list[dict[str, Any]]:
    text = _decode_text(content)
    preferred = "\t" if file_format == "tsv" else None
    candidates = [preferred] if preferred else []
    candidates.extend([";", ",", "\t", "|"])
    best: tuple[list[str], list[list[str]]] | None = None
    for delimiter in dict.fromkeys(item for item in candidates if item):
        try:
            reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            rows = []
            for row in reader:
                if len(rows) > MAX_PRODUCTS + 20:
                    raise CatalogFileLimitError("Il listino contiene troppi prodotti.")
                if len(row) > MAX_COLUMNS:
                    raise CatalogFileLimitError("Il listino contiene troppe colonne.")
                rows.append(row)
        except csv.Error:
            continue
        if rows and (best is None or len(rows[0]) > len(best[0])):
            best = (rows[0], rows[1:])
    if best is None or len(best[0]) < 1:
        raise CatalogFileValidationError("Formato CSV non riconosciuto.")
    return _normalized_rows(*best, source_row_start=2)


def _guard_xlsx(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise CatalogFileLimitError("Il file XLSX contiene troppi elementi.")
            total = 0
            for entry in entries:
                path = PurePath(entry.filename.replace("\\", "/"))
                if path.is_absolute() or ".." in path.parts or entry.flag_bits & 0x1:
                    raise CatalogFileValidationError("Il contenuto del file XLSX non è valido.")
                total += entry.file_size
                if total > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise CatalogFileLimitError("Il file XLSX decompresso è troppo grande.")
    except zipfile.BadZipFile as exc:
        raise CatalogFileValidationError("Il file XLSX non è valido.") from exc


def _best_header(rows: list[list[Any]]) -> tuple[int, list[Any]]:
    best_index = 0
    best_headers = rows[0] if rows else []
    best_score = -1
    for index, row in enumerate(rows[:10]):
        tokens = {_token(value) for value in row}
        score = sum(bool(tokens & aliases) for aliases in ALIAS_TOKENS.values())
        if score > best_score:
            best_index, best_headers, best_score = index, row, score
    return best_index, best_headers


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    _guard_xlsx(content)
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        raw: list[list[Any]] = []
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            if row_number > MAX_PRODUCTS + 10:
                raise CatalogFileLimitError("Il listino contiene troppi prodotti.")
            if len(cells) > MAX_COLUMNS:
                raise CatalogFileLimitError("Il listino contiene troppe colonne.")
            raw.append([_excel_text(cell.value, cell.number_format) for cell in cells])
            if len(raw) * max(1, len(cells)) > MAX_CELLS:
                raise CatalogFileLimitError("Il listino contiene troppe celle.")
    except (CatalogFileLimitError, CatalogFileValidationError):
        raise
    except Exception as exc:
        raise CatalogFileValidationError("Il file XLSX non è leggibile.") from exc
    finally:
        if workbook is not None:
            workbook.close()
    if not raw:
        raise CatalogFileValidationError("Il file XLSX non contiene fogli leggibili.")
    header_index, headers = _best_header(raw)
    return _normalized_rows(
        headers,
        raw[header_index + 1 :],
        source_row_start=header_index + 2,
    )


def _xls_cell_text(book, sheet, row: int, column: int) -> str:
    cell = sheet.cell(row, column)
    number_format = ""
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        try:
            xf = book.xf_list[cell.xf_index]
            number_format = book.format_map[xf.format_key].format_str
        except (AttributeError, IndexError, KeyError):
            number_format = ""
    return _excel_text(cell.value, number_format)


def _parse_xls(content: bytes) -> list[dict[str, Any]]:
    try:
        book = xlrd.open_workbook(file_contents=content, on_demand=True, formatting_info=True)
        sheet = book.sheet_by_index(0)
        if sheet.nrows > MAX_PRODUCTS + 10 or sheet.ncols > MAX_COLUMNS:
            raise CatalogFileLimitError("Il listino Excel supera i limiti consentiti.")
        if sheet.nrows * max(1, sheet.ncols) > MAX_CELLS:
            raise CatalogFileLimitError("Il listino contiene troppe celle.")
        raw = [
            [_xls_cell_text(book, sheet, row, column) for column in range(sheet.ncols)]
            for row in range(sheet.nrows)
        ]
    except (CatalogFileLimitError, CatalogFileValidationError):
        raise
    except Exception as exc:
        raise CatalogFileValidationError("Il file XLS non è leggibile.") from exc
    finally:
        try:
            book.release_resources()
        except (AttributeError, UnboundLocalError):
            pass
    if not raw:
        raise CatalogFileValidationError("Il file XLS non contiene fogli leggibili.")
    header_index, headers = _best_header(raw)
    return _normalized_rows(
        headers,
        raw[header_index + 1 :],
        source_row_start=header_index + 2,
    )


def _local_name(tag: Any) -> str:
    return str(tag or "").rsplit("}", 1)[-1].split(":")[-1]


def _flatten_xml(element) -> dict[str, str]:
    result: dict[str, str] = {}
    for child in element.iter():
        if child is element or list(child):
            continue
        key = _local_name(child.tag)
        value = "".join(child.itertext()).strip()
        if key and value and key not in result:
            result[key] = value
    return result


def _parse_xml(content: bytes) -> list[dict[str, Any]]:
    try:
        root = SafeElementTree.fromstring(content)
    except (DefusedXmlException, SafeElementTree.ParseError, ValueError) as exc:
        raise CatalogFileValidationError("Il file XML non è valido.") from exc
    count = 0
    for _element in root.iter():
        count += 1
        if count > MAX_XML_ELEMENTS:
            raise CatalogFileLimitError("Il file XML contiene troppi elementi.")
    named = [
        element for element in root.iter()
        if _local_name(element.tag).casefold() in {"product", "item", "offer", "record"}
        and list(element)
    ]
    candidates = named or [element for element in list(root) if list(element)]
    if not candidates and list(root):
        candidates = [root]
    records = [_flatten_xml(element) for element in candidates]
    records = [record for record in records if record]
    if not records:
        raise CatalogFileValidationError("Il file XML non contiene prodotti leggibili.")
    headers = list(dict.fromkeys(key for record in records for key in record))
    values = [[record.get(header, "") for header in headers] for record in records]
    return _normalized_rows(headers, values, source_row_start=2)


def parse_catalog(file_name: str, content: bytes) -> tuple[str, list[dict[str, Any]]]:
    if not content:
        raise CatalogFileValidationError("Il file del listino è vuoto.")
    if len(content) > MAX_CATALOG_ARTIFACT_BYTES:
        raise CatalogFileLimitError("Il file supera il limite di 20 MiB.")
    if not file_name or "\x00" in file_name or "/" in file_name or "\\" in file_name:
        raise CatalogFileValidationError("Il nome del file non è valido.")
    if len(file_name) > 255:
        raise CatalogFileValidationError("Il nome del file è troppo lungo.")
    suffix = PurePath(file_name).suffix.casefold().lstrip(".")
    if suffix in {"pkl", "pickle"}:
        raise CatalogFileValidationError("I file PKL/Pickle non sono accettati per sicurezza.")
    if suffix not in {"csv", "txt", "tsv", "xls", "xlsx", "xml"}:
        raise CatalogFileValidationError("Formato listino non supportato.")
    if suffix == "xlsx" and not content.startswith(b"PK"):
        raise CatalogFileValidationError("Il contenuto non corrisponde a un file XLSX.")
    if suffix == "xls" and not content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise CatalogFileValidationError("Il contenuto non corrisponde a un file XLS.")
    if suffix == "xml" and not content.lstrip().startswith((b"<", b"\xef\xbb\xbf<")):
        raise CatalogFileValidationError("Il contenuto non corrisponde a un file XML.")
    if suffix in {"csv", "txt", "tsv"}:
        rows = _parse_delimited(content, suffix)
    elif suffix == "xlsx":
        rows = _parse_xlsx(content)
    elif suffix == "xls":
        rows = _parse_xls(content)
    else:
        rows = _parse_xml(content)
    return suffix, rows
