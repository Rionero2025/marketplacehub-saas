"""Safe Kaufland Seller Portal tracking import primitives.

Files are parsed in memory and deliberately never archived by this module.
"""

from __future__ import annotations

import csv
import io
import posixpath
import re
import unicodedata
import zipfile
import zlib
from pathlib import Path
from xml.sax import SAXException
from xml.sax.handler import ContentHandler, feature_namespaces

from defusedxml import sax as defused_sax

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_XLSX_BYTES = 50 * 1024 * 1024
MAX_XLSX_ENTRIES = 1_000
MAX_XLSX_SHARED_STRINGS_BYTES = 10 * 1024 * 1024
MAX_XLSX_SHARED_STRINGS = 100_000
MAX_XLSX_SHARED_STRING_LENGTH = 2_000
MAX_XLSX_SHARED_STRING_ELEMENTS = 256
MAX_XLSX_SHARED_TEXT = 5 * 1024 * 1024
MAX_XLSX_SHARED_ELEMENTS = 500_000
MAX_XLSX_STYLES_BYTES = 2 * 1024 * 1024
MAX_XLSX_STYLE_ELEMENTS = 10_000
MAX_XLSX_GLOBAL_PART_BYTES = 2 * 1024 * 1024
MAX_XLSX_GLOBAL_ELEMENTS = 10_000
MAX_XLSX_GLOBAL_TEXT = 2 * 1024 * 1024
MAX_XLSX_DEFINED_NAMES = 2_000
MAX_XLSX_DEFINED_NAME_LENGTH = 2_000
MAX_XLSX_WORKSHEET_ELEMENTS = 4_100_000
MAX_XLSX_WORKSHEET_TEXT = 5 * 1024 * 1024
MAX_XLSX_CELL_ELEMENTS = 256
MAX_XLSX_WORKSHEET_PREAMBLE_ELEMENTS = 10_000
MAX_XLSX_WORKSHEET_PREAMBLE_TEXT = 256 * 1024
MAX_XLSX_XML_ATTRIBUTE_BYTES = 8 * 1024
MAX_XLSX_XML_TAG_BYTES = 64 * 1024
MAX_XLSX_XML_ELEMENTS_TOTAL = 1_000_000
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_CELLS = 200_000
MAX_HEADER_LENGTH = 200
MAX_IDENTIFIER_LENGTH = 200
MAX_CELL_LENGTH = 2_000
MAX_TRACKING_TARGET_UNITS = 10_000
PREVIEW_ROWS = 10
FORMATS = ("csv", "xlsx", "xls")
MAPPING_FIELDS = (
    "id_order_unit", "id_order", "carrier_code", "tracking_numbers", "combined_shipment",
)

TRACKING_IMPORT_ALIASES = {
    "id_order_unit": {
        "idorderunit", "orderunitid", "orderunit", "unitaordine", "unitaordineid",
        "idunitaordine", "bestellpositionid", "orderitemid",
    },
    "id_order": {
        "idorder", "orderid", "ordernumber", "ordine", "numeroordine", "bestellnummer",
        "bestellung", "kauflandorderid",
    },
    "carrier_code": {
        "carriercode", "carrier", "carriername", "corriere", "speditocon",
        "versanddienstleister", "shippingprovider", "shippingcarrier",
    },
    "tracking_numbers": {
        "trackingnumbers", "trackingnumber", "tracking", "trackingcode", "numerotracking",
        "numerospedizione", "tracciabilita", "sendungsnummer", "paketnummer", "parcelnumber",
    },
    "combined_shipment": {
        "shipment", "shipmentinformation", "shippinginformation", "informazionispedizione",
        "spedizione", "versandinformation",
    },
}


class TrackingValidationError(ValueError):
    pass


class TrackingLimitError(TrackingValidationError):
    pass


_SPREADSHEETML_NAMESPACES = {
    "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "http://purl.oclc.org/ooxml/spreadsheetml/main",
}
_PACKAGE_RELATIONSHIPS_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
_OFFICE_RELATIONSHIPS_NAMESPACES = {
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "http://purl.oclc.org/ooxml/officeDocument/relationships",
}
_CELL_REFERENCE = re.compile(r"^([A-Za-z]{1,3})([1-9][0-9]*)$")
_DIMENSION_REFERENCE = re.compile(
    r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)"
    r"(?::\$?([A-Za-z]{1,3})\$?([1-9][0-9]*))?$"
)
_XML_TAG_BOUNDARY = re.compile(br"[\"'>]")


class _XmlLexicalGuard:
    """Bound XML markup tokens before Expat materializes attribute strings."""

    def __init__(self):
        self.mode = "text"
        self.open_prefix = bytearray()
        self.quote: int | None = None
        self.token_bytes = 0
        self.attribute_bytes = 0
        self.section_tail = b""

    def _start_token(self) -> None:
        self.mode = "open"
        self.open_prefix = bytearray(b"<")
        self.quote = None
        self.token_bytes = 1
        self.attribute_bytes = 0
        self.section_tail = b""

    def _finish_token(self) -> None:
        self.mode = "text"
        self.open_prefix.clear()
        self.quote = None
        self.token_bytes = 0
        self.attribute_bytes = 0
        self.section_tail = b""

    def _check_token_size(self) -> None:
        if self.token_bytes > MAX_XLSX_XML_TAG_BYTES:
            raise TrackingLimitError(
                "Il file XLSX contiene un elemento XML troppo grande."
            )

    def _increase_token(self, amount: int) -> None:
        self.token_bytes += amount
        self._check_token_size()

    def _feed_quoted_markup(self, value: int) -> None:
        if self.quote is not None:
            if value == self.quote:
                self.quote = None
                self.attribute_bytes = 0
            else:
                self.attribute_bytes += 1
                if self.attribute_bytes > MAX_XLSX_XML_ATTRIBUTE_BYTES:
                    raise TrackingLimitError(
                        "Il file XLSX contiene un attributo XML troppo grande."
                    )
        elif value in {ord('"'), ord("'")}:
            self.quote = value
            self.attribute_bytes = 0
        elif value == ord(">"):
            self._finish_token()

    def _feed_open(self, value: int) -> None:
        self.open_prefix.append(value)
        prefix = bytes(self.open_prefix)
        if prefix == b"<!--":
            self.mode = "comment"
            self.section_tail = b""
        elif prefix == b"<![CDATA[":
            self.mode = "cdata"
            self.section_tail = b""
        elif prefix == b"<?":
            self.mode = "processing"
            self.section_tail = b""
        elif any(candidate.startswith(prefix) for candidate in (b"<!--", b"<![CDATA[", b"<?")):
            return
        else:
            self.mode = "declaration" if prefix.startswith(b"<!") else "tag"
            buffered = tuple(self.open_prefix[1:])
            self.open_prefix.clear()
            for buffered_value in buffered:
                self._feed_quoted_markup(buffered_value)

    def feed(self, chunk: bytes) -> None:
        index = 0
        size = len(chunk)
        while index < size:
            if self.mode == "text":
                opening = chunk.find(b"<", index)
                if opening < 0:
                    return
                self._start_token()
                index = opening + 1
                continue
            if self.mode == "open":
                value = chunk[index]
                index += 1
                self._increase_token(1)
                self._feed_open(value)
                continue
            if self.mode in {"comment", "cdata", "processing"}:
                terminator = {
                    "comment": b"-->", "cdata": b"]]>", "processing": b"?>",
                }[self.mode]
                remaining = chunk[index:]
                combined = self.section_tail + remaining
                closing = combined.find(terminator)
                if closing < 0:
                    self._increase_token(len(remaining))
                    self.section_tail = combined[-(len(terminator) - 1):]
                    return
                consumed = max(0, closing + len(terminator) - len(self.section_tail))
                self._increase_token(consumed)
                index += consumed
                self._finish_token()
                continue
            if self.quote is not None:
                closing = chunk.find(bytes((self.quote,)), index)
                consumed = size - index if closing < 0 else closing - index
                self._increase_token(consumed)
                self.attribute_bytes += consumed
                if self.attribute_bytes > MAX_XLSX_XML_ATTRIBUTE_BYTES:
                    raise TrackingLimitError(
                        "Il file XLSX contiene un attributo XML troppo grande."
                    )
                index += consumed
                if closing < 0:
                    return
                self._increase_token(1)
                index += 1
                self.quote = None
                self.attribute_bytes = 0
                continue
            match = _XML_TAG_BOUNDARY.search(chunk, index)
            if match is None:
                self._increase_token(size - index)
                return
            boundary = match.start()
            self._increase_token(boundary - index + 1)
            boundary_value = chunk[boundary]
            index = boundary + 1
            if boundary_value in {ord('"'), ord("'")}:
                self.quote = boundary_value
                self.attribute_bytes = 0
            else:
                self._finish_token()


class _XmlBudget:
    """Shared complexity budget across every XML part parsed by the guard."""

    def __init__(self):
        self.elements = 0

    def add_element(self) -> None:
        self.elements += 1
        if self.elements > MAX_XLSX_XML_ELEMENTS_TOTAL:
            raise TrackingLimitError(
                "Il file XLSX contiene troppi elementi XML complessivi."
            )


class _BudgetedContentHandler(ContentHandler):
    def __init__(self, delegate: ContentHandler, budget: _XmlBudget):
        super().__init__()
        self.delegate = delegate
        self.budget = budget

    def startElementNS(self, name, qname, attrs):  # noqa: N802
        self.budget.add_element()
        self.delegate.startElementNS(name, qname, attrs)

    def endElementNS(self, name, qname):  # noqa: N802
        self.delegate.endElementNS(name, qname)

    def characters(self, content):
        self.delegate.characters(content)


class _SharedStringsGuard(ContentHandler):
    """Count the global XLSX string table without materializing its elements."""

    def __init__(self):
        super().__init__()
        self.root_seen = False
        self.element_count = 0
        self.string_count = 0
        self.current_string_length: int | None = None
        self.current_string_elements = 0
        self.total_text = 0

    def startElementNS(self, name, _qname, _attrs):  # noqa: N802
        namespace, local_name = name
        if not self.root_seen:
            if local_name != "sst" or namespace not in _SPREADSHEETML_NAMESPACES:
                raise TrackingValidationError(
                    "La tabella delle stringhe XLSX non è valida."
                )
            self.root_seen = True
        self.element_count += 1
        if self.element_count > MAX_XLSX_SHARED_ELEMENTS:
            raise TrackingLimitError(
                "La tabella delle stringhe XLSX contiene troppi elementi."
            )
        if self.current_string_length is not None:
            self.current_string_elements += 1
            if self.current_string_elements > MAX_XLSX_SHARED_STRING_ELEMENTS:
                raise TrackingLimitError(
                    "Una stringa XLSX contiene troppi elementi."
                )
        if local_name == "si" and namespace in _SPREADSHEETML_NAMESPACES:
            if self.current_string_length is not None:
                raise TrackingValidationError(
                    "La tabella delle stringhe XLSX non è valida."
                )
            self.string_count += 1
            if self.string_count > MAX_XLSX_SHARED_STRINGS:
                raise TrackingLimitError(
                    "La tabella delle stringhe XLSX supera il limite consentito."
                )
            self.current_string_length = 0
            self.current_string_elements = 1

    def characters(self, content):
        if self.current_string_length is None:
            return
        length = len(content)
        self.current_string_length += length
        self.total_text += length
        if self.current_string_length > MAX_XLSX_SHARED_STRING_LENGTH:
            raise TrackingLimitError(
                "La tabella delle stringhe XLSX contiene un valore troppo grande."
            )
        if self.total_text > MAX_XLSX_SHARED_TEXT:
            raise TrackingLimitError(
                "La tabella delle stringhe XLSX contiene troppi dati."
            )

    def endElementNS(self, name, _qname):  # noqa: N802
        namespace, local_name = name
        if local_name == "si" and namespace in _SPREADSHEETML_NAMESPACES:
            self.current_string_length = None
            self.current_string_elements = 0


class _StylesGuard(ContentHandler):
    """Bound XLSX stylesheet complexity before openpyxl builds style objects."""

    def __init__(self):
        super().__init__()
        self.root_seen = False
        self.element_count = 0

    def startElementNS(self, name, _qname, _attrs):  # noqa: N802
        namespace, local_name = name
        if not self.root_seen:
            if local_name != "styleSheet" or namespace not in _SPREADSHEETML_NAMESPACES:
                raise TrackingValidationError("Il foglio di stile XLSX non è valido.")
            self.root_seen = True
        self.element_count += 1
        if self.element_count > MAX_XLSX_STYLE_ELEMENTS:
            raise TrackingLimitError(
                "Il foglio di stile XLSX contiene troppi elementi."
            )


class _GlobalXmlGuard(ContentHandler):
    """Bound OOXML parts which openpyxl builds as in-memory object trees."""

    def __init__(self):
        super().__init__()
        self.element_count = 0
        self.text_length = 0

    def startElementNS(self, _name, _qname, _attrs):  # noqa: N802
        self.element_count += 1
        if self.element_count > MAX_XLSX_GLOBAL_ELEMENTS:
            raise TrackingLimitError(
                "Una parte globale del file XLSX contiene troppi elementi."
            )

    def characters(self, content):
        self.text_length += len(content)
        if self.text_length > MAX_XLSX_GLOBAL_TEXT:
            raise TrackingLimitError(
                "Una parte globale del file XLSX contiene troppi dati."
            )


class _WorkbookGuard(_GlobalXmlGuard):
    """Apply stricter limits to names eagerly materialized from workbook.xml."""

    def __init__(self):
        super().__init__()
        self.root_seen = False
        self.defined_name_count = 0
        self.current_defined_name_length: int | None = None
        self.sheet_relationship_ids: list[str] = []
        self.in_sheets = False
        self.sheets_seen = False

    def startElementNS(self, name, qname, attrs):  # noqa: N802
        namespace, local_name = name
        if not self.root_seen:
            if local_name != "workbook" or namespace not in _SPREADSHEETML_NAMESPACES:
                raise TrackingValidationError("La definizione del file XLSX non è valida.")
            self.root_seen = True
        super().startElementNS(name, qname, attrs)
        if local_name == "definedName" and namespace in _SPREADSHEETML_NAMESPACES:
            if self.current_defined_name_length is not None:
                raise TrackingValidationError(
                    "La definizione dei nomi XLSX non è valida."
                )
            self.defined_name_count += 1
            if self.defined_name_count > MAX_XLSX_DEFINED_NAMES:
                raise TrackingLimitError(
                    "Il file XLSX contiene troppi nomi definiti."
                )
            self.current_defined_name_length = 0
        elif local_name == "sheets" and namespace in _SPREADSHEETML_NAMESPACES:
            if self.sheets_seen:
                raise TrackingValidationError(
                    "La definizione dei fogli XLSX non è valida."
                )
            self.sheets_seen = True
            self.in_sheets = True
        elif (
            local_name == "sheet"
            and namespace in _SPREADSHEETML_NAMESPACES
            and self.in_sheets
        ):
            relationship_id = next(
                (
                    str(attrs[(relationship_namespace, "id")])
                    for relationship_namespace in _OFFICE_RELATIONSHIPS_NAMESPACES
                    if (relationship_namespace, "id") in attrs
                ),
                "",
            )
            if not relationship_id:
                raise TrackingValidationError(
                    "La relazione di un foglio XLSX non è valida."
                )
            self.sheet_relationship_ids.append(relationship_id)

    def characters(self, content):
        super().characters(content)
        if self.current_defined_name_length is None:
            return
        self.current_defined_name_length += len(content)
        if self.current_defined_name_length > MAX_XLSX_DEFINED_NAME_LENGTH:
            raise TrackingLimitError(
                "Il file XLSX contiene un nome definito troppo grande."
            )

    def endElementNS(self, name, _qname):  # noqa: N802
        namespace, local_name = name
        if local_name == "definedName" and namespace in _SPREADSHEETML_NAMESPACES:
            self.current_defined_name_length = None
        elif local_name == "sheets" and namespace in _SPREADSHEETML_NAMESPACES:
            self.in_sheets = False


class _WorkbookRelationshipsGuard(_GlobalXmlGuard):
    """Find every worksheet target while bounding the relationship table."""

    def __init__(self):
        super().__init__()
        self.root_seen = False
        self.worksheet_names: set[str] = set()
        self.worksheets_by_id: dict[str, str] = {}

    def startElementNS(self, name, qname, attrs):  # noqa: N802
        namespace, local_name = name
        if not self.root_seen:
            if (
                local_name != "Relationships"
                or namespace != _PACKAGE_RELATIONSHIPS_NAMESPACE
            ):
                raise TrackingValidationError(
                    "Le relazioni del file XLSX non sono valide."
                )
            self.root_seen = True
        super().startElementNS(name, qname, attrs)
        if local_name != "Relationship" or namespace != _PACKAGE_RELATIONSHIPS_NAMESPACE:
            return
        relationship_type = str(attrs.get((None, "Type"), ""))
        if relationship_type.rsplit("/", 1)[-1] != "worksheet":
            return
        if str(attrs.get((None, "TargetMode"), "")).casefold() == "external":
            raise TrackingValidationError("Un foglio XLSX non può essere esterno.")
        target = str(attrs.get((None, "Target"), ""))
        relationship_id = str(attrs.get((None, "Id"), ""))
        if not relationship_id or not target or "\\" in target or "\x00" in target:
            raise TrackingValidationError("Il percorso di un foglio XLSX non è valido.")
        resolved = (
            posixpath.normpath(target.lstrip("/"))
            if target.startswith("/")
            else posixpath.normpath(posixpath.join("xl", target))
        )
        if resolved in {"", ".", ".."} or resolved.startswith("../"):
            raise TrackingValidationError("Il percorso di un foglio XLSX non è valido.")
        if relationship_id in self.worksheets_by_id:
            raise TrackingValidationError("Le relazioni dei fogli XLSX sono duplicate.")
        self.worksheet_names.add(resolved)
        self.worksheets_by_id[relationship_id] = resolved


def _column_number(letters: str) -> int:
    result = 0
    for character in letters.upper():
        result = (result * 26) + ord(character) - ord("A") + 1
    return result


class _WorksheetGuard(ContentHandler):
    """Validate worksheet shape before openpyxl expands a logical row in memory."""

    def __init__(self):
        super().__init__()
        self.root_seen = False
        self.in_sheet_data = False
        self.row_count = 0
        self.last_row_number = 0
        self.max_row_number = 0
        self.max_column_number = 0
        self.current_row: int | None = None
        self.current_row_cells = 0
        self.cell_count = 0
        self.last_column = 0
        self.current_cell_text: int | None = None
        self.current_cell_elements = 0
        self.element_count = 0
        self.text_length = 0

    def _check_logical_grid(self) -> None:
        if self.max_row_number * self.max_column_number > MAX_CELLS:
            raise TrackingLimitError(f"Il file supera il limite di {MAX_CELLS} celle.")

    def startElementNS(self, name, _qname, attrs):  # noqa: N802
        namespace, local_name = name
        if not self.root_seen:
            if local_name != "worksheet" or namespace not in _SPREADSHEETML_NAMESPACES:
                raise TrackingValidationError("La struttura di un foglio XLSX non è valida.")
            self.root_seen = True
        self.element_count += 1
        if self.element_count > MAX_XLSX_WORKSHEET_ELEMENTS:
            raise TrackingLimitError("Un foglio XLSX contiene troppi elementi.")
        if self.current_cell_text is not None:
            self.current_cell_elements += 1
            if self.current_cell_elements > MAX_XLSX_CELL_ELEMENTS:
                raise TrackingLimitError("Una cella XLSX contiene troppi elementi.")
        if namespace not in _SPREADSHEETML_NAMESPACES:
            return
        if local_name == "dimension":
            reference = str(attrs.get((None, "ref"), ""))
            match = _DIMENSION_REFERENCE.fullmatch(reference)
            if match is None:
                raise TrackingValidationError("La dimensione del foglio XLSX non è valida.")
            minimum_column = _column_number(match.group(1))
            minimum_row = int(match.group(2))
            maximum_column = _column_number(match.group(3) or match.group(1))
            maximum_row = int(match.group(4) or match.group(2))
            if minimum_column > maximum_column or minimum_row > maximum_row:
                raise TrackingValidationError("La dimensione del foglio XLSX non è valida.")
            if maximum_column > MAX_COLUMNS:
                raise TrackingLimitError(f"Il file supera il limite di {MAX_COLUMNS} colonne.")
            if maximum_row > MAX_ROWS + 1:
                raise TrackingLimitError(f"Il file supera il limite di {MAX_ROWS} righe.")
            self.max_column_number = max(self.max_column_number, maximum_column)
            self.max_row_number = max(self.max_row_number, maximum_row)
            self._check_logical_grid()
        elif local_name == "sheetData":
            if self.in_sheet_data:
                raise TrackingValidationError("La struttura di un foglio XLSX non è valida.")
            self.in_sheet_data = True
        elif local_name == "row" and self.in_sheet_data:
            if self.current_row is not None:
                raise TrackingValidationError("La struttura delle righe XLSX non è valida.")
            raw_number = str(attrs.get((None, "r"), ""))
            if raw_number:
                if not raw_number.isascii() or not raw_number.isdecimal():
                    raise TrackingValidationError(
                        "La coordinata di una riga XLSX non è valida."
                    )
                row_number = int(raw_number)
            else:
                row_number = self.last_row_number + 1
            if row_number <= self.last_row_number:
                raise TrackingValidationError(
                    "Le coordinate delle righe XLSX sono duplicate o non ordinate."
                )
            if row_number > MAX_ROWS + 1:
                raise TrackingLimitError(f"Il file supera il limite di {MAX_ROWS} righe.")
            self.max_row_number = max(self.max_row_number, row_number)
            self._check_logical_grid()
            self.row_count += 1
            if self.row_count > MAX_ROWS + 1:
                raise TrackingLimitError(f"Il file supera il limite di {MAX_ROWS} righe.")
            self.current_row = row_number
            self.current_row_cells = 0
            self.last_column = 0
        elif local_name == "c" and self.in_sheet_data:
            if self.current_row is None or self.current_cell_text is not None:
                raise TrackingValidationError("La struttura delle celle XLSX non è valida.")
            reference = str(attrs.get((None, "r"), ""))
            if reference:
                match = _CELL_REFERENCE.fullmatch(reference)
                if match is None:
                    raise TrackingValidationError(
                        "La coordinata di una cella XLSX non è valida."
                    )
                column_number = _column_number(match.group(1))
                row_number = int(match.group(2))
                if row_number != self.current_row:
                    raise TrackingValidationError(
                        "La coordinata di una cella XLSX non è valida."
                    )
            else:
                column_number = self.last_column + 1
            if column_number <= self.last_column:
                raise TrackingValidationError(
                    "Le coordinate delle celle XLSX sono duplicate o non ordinate."
                )
            if column_number > MAX_COLUMNS:
                raise TrackingLimitError(f"Il file supera il limite di {MAX_COLUMNS} colonne.")
            self.max_column_number = max(self.max_column_number, column_number)
            self._check_logical_grid()
            self.current_row_cells += 1
            if self.current_row_cells > MAX_COLUMNS:
                raise TrackingLimitError(f"Il file supera il limite di {MAX_COLUMNS} colonne.")
            self.cell_count += 1
            if self.cell_count > MAX_CELLS:
                raise TrackingLimitError(f"Il file supera il limite di {MAX_CELLS} celle.")
            self.last_column = column_number
            self.current_cell_text = 0
            self.current_cell_elements = 1

    def characters(self, content):
        length = len(content)
        self.text_length += length
        if self.text_length > MAX_XLSX_WORKSHEET_TEXT:
            raise TrackingLimitError("Un foglio XLSX contiene troppi dati.")
        if self.current_cell_text is None:
            return
        self.current_cell_text += length
        if self.current_cell_text > MAX_CELL_LENGTH:
            raise TrackingLimitError("Il file contiene una cella troppo grande.")

    def endElementNS(self, name, _qname):  # noqa: N802
        namespace, local_name = name
        if namespace not in _SPREADSHEETML_NAMESPACES:
            return
        if local_name == "c" and self.in_sheet_data:
            if self.current_cell_text is None:
                raise TrackingValidationError("La struttura delle celle XLSX non è valida.")
            self.current_cell_text = None
            self.current_cell_elements = 0
        elif local_name == "row" and self.in_sheet_data:
            if self.current_row is None or self.current_cell_text is not None:
                raise TrackingValidationError("La struttura delle righe XLSX non è valida.")
            self.last_row_number = self.current_row
            self.current_row = None
        elif local_name == "sheetData":
            if self.current_row is not None or self.current_cell_text is not None:
                raise TrackingValidationError("La struttura di un foglio XLSX non è valida.")
            self.in_sheet_data = False


class _StopXmlScan(Exception):
    pass


class _UnselectedWorksheetGuard(ContentHandler):
    """Bound only the preamble parsed by ReadOnlyWorksheet._get_size()."""

    def __init__(self):
        super().__init__()
        self.root_seen = False
        self.element_count = 0
        self.text_length = 0

    def startElementNS(self, name, _qname, _attrs):  # noqa: N802
        namespace, local_name = name
        if not self.root_seen:
            if local_name != "worksheet" or namespace not in _SPREADSHEETML_NAMESPACES:
                raise TrackingValidationError("La struttura di un foglio XLSX non è valida.")
            self.root_seen = True
        if (
            namespace in _SPREADSHEETML_NAMESPACES
            and local_name in {"dimension", "sheetData"}
        ):
            raise _StopXmlScan
        self.element_count += 1
        if self.element_count > MAX_XLSX_WORKSHEET_PREAMBLE_ELEMENTS:
            raise TrackingLimitError(
                "Il preambolo di un foglio XLSX contiene troppi elementi."
            )

    def characters(self, content):
        self.text_length += len(content)
        if self.text_length > MAX_XLSX_WORKSHEET_PREAMBLE_TEXT:
            raise TrackingLimitError(
                "Il preambolo di un foglio XLSX contiene troppi dati."
            )


def _stream_guard_xlsx_xml(
    source, content_handler: ContentHandler, budget: _XmlBudget,
) -> None:
    parser = defused_sax.make_parser()
    parser.setFeature(feature_namespaces, True)
    parser.setContentHandler(_BudgetedContentHandler(content_handler, budget))
    parser.forbid_dtd = True
    parser.forbid_entities = True
    parser.forbid_external = True
    lexical_guard = _XmlLexicalGuard()
    try:
        while chunk := source.read(64 * 1024):
            # Check markup first so Expat never receives an oversized token.
            lexical_guard.feed(chunk)
            parser.feed(chunk)
        parser.close()
    except _StopXmlScan:
        return
    except TrackingValidationError:
        raise
    except (SAXException, ValueError, UnicodeError) as exc:
        raise TrackingValidationError("Il file XLSX non è valido.") from exc


def _is_streamed_worksheet_part(file_name: str) -> bool:
    normalized = file_name.casefold()
    return (
        normalized.startswith("xl/worksheets/")
        and "/" not in normalized[len("xl/worksheets/"):]
        and normalized.endswith(".xml")
    )


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        result = "true" if value else "false"
    elif isinstance(value, float) and value.is_integer():
        result = str(int(value))
    else:
        result = str(value)
    if "\x00" in result:
        raise TrackingValidationError("Il valore contiene caratteri non validi.")
    return result.strip()[:MAX_CELL_LENGTH]


def normalized_header(value) -> str:
    text = unicodedata.normalize("NFKD", clean(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def detect_tracking_columns(columns) -> dict[str, str]:
    result: dict[str, str] = {}
    aliases = {
        field: {normalized_header(alias) for alias in values}
        for field, values in TRACKING_IMPORT_ALIASES.items()
    }
    for column in columns:
        normalized = normalized_header(column)
        for field, candidates in aliases.items():
            if field not in result and normalized in candidates:
                result[field] = str(column)
                break
    return result


def split_shipment_text(value) -> tuple[str, str]:
    text = clean(value)
    if not text:
        return "", ""
    for separator in ("|", " - ", ";"):
        if separator in text:
            carrier, tracking = text.split(separator, 1)
            return carrier.strip(), tracking.strip()
    return "", text


def normalize_tracking(value) -> str:
    parts = (part.strip() for part in clean(value).split(","))
    return ", ".join(dict.fromkeys(part for part in parts if part))


def validate_file(file_name: str, content: bytes) -> str:
    raw_name = str(file_name or "")
    if len(raw_name) > 255 or "\x00" in raw_name or "/" in raw_name or "\\" in raw_name:
        raise TrackingValidationError("Il nome del file non è valido.")
    name = Path(raw_name).name
    suffix = Path(name).suffix.lower().lstrip(".")
    if not name or suffix not in FORMATS:
        raise TrackingValidationError("Formato non supportato. Usa CSV, XLSX oppure XLS.")
    if not content:
        raise TrackingValidationError("Il file è vuoto.")
    if len(content) > MAX_FILE_BYTES:
        raise TrackingLimitError("Il file supera il limite di 5 MB.")
    if suffix == "csv" and b"\x00" in content:
        raise TrackingValidationError("Il CSV contiene caratteri non validi.")
    return suffix


def _unique_columns(values) -> list[str]:
    raw_columns = ["" if value is None else str(value).strip() for value in values]
    if any(len(value) > MAX_HEADER_LENGTH for value in raw_columns):
        raise TrackingLimitError("Il file contiene un'intestazione troppo grande.")
    columns = [clean(value) for value in raw_columns]
    if not columns or not any(columns):
        raise TrackingValidationError("Il file non contiene intestazioni valide.")
    if len(columns) > MAX_COLUMNS:
        raise TrackingLimitError(f"Il file supera il limite di {MAX_COLUMNS} colonne.")
    if any(not column for column in columns):
        raise TrackingValidationError("Tutte le colonne devono avere un'intestazione.")
    if len({column.casefold() for column in columns}) != len(columns):
        raise TrackingValidationError("Il file contiene intestazioni duplicate.")
    return columns


def _bounded_rows(columns: list[str], values) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    logical_cells = len(columns)
    for raw in values:
        raw_cells = list(raw)
        # csv.reader returns [] only for a physically blank line. A delimited
        # row (for example ";") is a real source row and must be reported invalid.
        if not raw_cells:
            continue
        if len(raw_cells) > len(columns):
            raise TrackingValidationError("Una riga contiene più celle delle intestazioni.")
        if any(len(str(value)) > MAX_CELL_LENGTH for value in raw_cells if value is not None):
            raise TrackingLimitError("Il file contiene una cella troppo grande.")
        logical_cells += len(columns)
        if logical_cells > MAX_CELLS:
            raise TrackingLimitError(f"Il file supera il limite di {MAX_CELLS} celle.")
        cells = [clean(value) for value in raw_cells[: len(columns)]]
        cells.extend([""] * (len(columns) - len(cells)))
        rows.append(dict(zip(columns, cells, strict=True)))
        if len(rows) > MAX_ROWS:
            raise TrackingLimitError(f"Il file supera il limite di {MAX_ROWS} righe.")
    return rows


def _validate_csv_shape(text: str, dialect) -> None:
    """Reject oversized logical records before csv.reader allocates their cells."""
    delimiter = dialect.delimiter
    quotechar = None if dialect.quoting == csv.QUOTE_NONE else dialect.quotechar
    escapechar = dialect.escapechar
    doublequote = bool(dialect.doublequote)
    skipinitialspace = bool(dialect.skipinitialspace)
    fields = 1
    cell_length = 0
    in_quotes = False
    field_start = True
    index = 0
    while index < len(text):
        char = text[index]
        if escapechar and char == escapechar and index + 1 < len(text):
            cell_length += 1
            field_start = False
            index += 2
        elif in_quotes:
            if char == quotechar:
                if doublequote and index + 1 < len(text) and text[index + 1] == quotechar:
                    cell_length += 1
                    index += 2
                else:
                    in_quotes = False
                    index += 1
            else:
                cell_length += 1
                index += 1
        elif quotechar and char == quotechar and field_start:
            in_quotes = True
            field_start = False
            index += 1
        elif char == delimiter:
            fields += 1
            if fields > MAX_COLUMNS:
                raise TrackingLimitError(
                    f"Il file supera il limite di {MAX_COLUMNS} colonne."
                )
            cell_length = 0
            field_start = True
            index += 1
        elif char in "\r\n":
            if char == "\r" and index + 1 < len(text) and text[index + 1] == "\n":
                index += 1
            fields = 1
            cell_length = 0
            field_start = True
            index += 1
        elif skipinitialspace and field_start and char == " ":
            index += 1
        else:
            cell_length += 1
            field_start = False
            index += 1
        if cell_length > MAX_CELL_LENGTH:
            raise TrackingLimitError("Il CSV contiene una cella troppo grande.")


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    text = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise TrackingValidationError("La codifica del CSV non è riconosciuta.")
    sample = text[:32_768]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    _validate_csv_shape(text, dialect)
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        columns = _unique_columns(next(reader))
        return columns, _bounded_rows(columns, reader)
    except StopIteration as exc:
        raise TrackingValidationError("Il CSV è vuoto.") from exc
    except csv.Error as exc:
        error = str(exc).casefold()
        if "field larger" in error:
            raise TrackingLimitError("Il CSV contiene una cella troppo grande.") from exc
        raise TrackingValidationError("Il CSV non è leggibile.") from exc


def _guard_xlsx(content: bytes) -> None:
    forbidden = (b"<!DOCTYPE", b"<!ENTITY")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise TrackingLimitError("Il file XLSX contiene troppi elementi.")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_XLSX_BYTES:
                raise TrackingLimitError(
                    "Il file XLSX compresso supera i limiti di sicurezza."
                )
            if any(
                entry.flag_bits & 0x1
                or entry.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
                for entry in entries
            ):
                raise TrackingValidationError(
                    "Il file XLSX usa una compressione o cifratura non consentita."
                )
            if not any(entry.filename == "xl/workbook.xml" for entry in entries):
                raise TrackingValidationError("Il file XLSX non è valido.")
            expanded_bytes = 0
            for entry in entries:
                if not entry.filename.casefold().endswith((".xml", ".rels")):
                    continue
                carry = b""
                with archive.open(entry) as source:
                    while chunk := source.read(64 * 1024):
                        expanded_bytes += len(chunk)
                        if expanded_bytes > MAX_UNCOMPRESSED_XLSX_BYTES:
                            raise TrackingLimitError(
                                "Il file XLSX compresso supera i limiti di sicurezza."
                            )
                        scanned = (carry + chunk).upper()
                        if any(marker in scanned for marker in forbidden):
                            raise TrackingValidationError(
                                "Il file XLSX contiene dichiarazioni XML non consentite."
                            )
                        carry = scanned[-16:]
            entries_by_name = {entry.filename: entry for entry in entries}
            xml_budget = _XmlBudget()
            workbook_entry = entries_by_name["xl/workbook.xml"]
            if workbook_entry.file_size > MAX_XLSX_GLOBAL_PART_BYTES:
                raise TrackingLimitError("Una parte globale del file XLSX è troppo grande.")
            workbook_guard = _WorkbookGuard()
            with archive.open(workbook_entry) as source:
                _stream_guard_xlsx_xml(source, workbook_guard, xml_budget)
            relationships_guard = _WorkbookRelationshipsGuard()
            workbook_relationships = entries_by_name.get("xl/_rels/workbook.xml.rels")
            if workbook_relationships is not None:
                if workbook_relationships.file_size > MAX_XLSX_GLOBAL_PART_BYTES:
                    raise TrackingLimitError(
                        "Una parte globale del file XLSX è troppo grande."
                    )
                with archive.open(workbook_relationships) as source:
                    _stream_guard_xlsx_xml(source, relationships_guard, xml_budget)
            worksheet_names = {
                entry.filename for entry in entries
                if _is_streamed_worksheet_part(entry.filename)
            } | relationships_guard.worksheet_names
            first_worksheet_name = next(
                (
                    relationships_guard.worksheets_by_id[relationship_id]
                    for relationship_id in workbook_guard.sheet_relationship_ids
                    if (
                        relationship_id in relationships_guard.worksheets_by_id
                        and relationships_guard.worksheets_by_id[relationship_id]
                        in entries_by_name
                    )
                ),
                None,
            )
            for entry in entries:
                normalized_name = entry.filename.casefold()
                if not normalized_name.endswith((".xml", ".rels")):
                    continue
                if entry.filename == first_worksheet_name:
                    maximum_bytes = MAX_UNCOMPRESSED_XLSX_BYTES
                    handler = _WorksheetGuard()
                elif entry.filename in worksheet_names:
                    # The original reads sheet_name=0. In read-only mode openpyxl
                    # only parses their preamble to discover worksheet dimensions.
                    maximum_bytes = MAX_UNCOMPRESSED_XLSX_BYTES
                    handler = _UnselectedWorksheetGuard()
                elif entry.filename == "xl/sharedStrings.xml":
                    maximum_bytes = MAX_XLSX_SHARED_STRINGS_BYTES
                    handler = _SharedStringsGuard()
                elif entry.filename == "xl/styles.xml":
                    maximum_bytes = MAX_XLSX_STYLES_BYTES
                    handler = _StylesGuard()
                elif entry.filename in {
                    "xl/workbook.xml", "xl/_rels/workbook.xml.rels",
                }:
                    continue
                else:
                    maximum_bytes = MAX_XLSX_GLOBAL_PART_BYTES
                    handler = _GlobalXmlGuard()
                if entry.file_size > maximum_bytes:
                    raise TrackingLimitError(
                        "Una parte globale del file XLSX è troppo grande."
                    )
                with archive.open(entry) as source:
                    _stream_guard_xlsx_xml(source, handler, xml_budget)
    except TrackingValidationError:
        raise
    except (
        zipfile.BadZipFile, RuntimeError, NotImplementedError, OSError, EOFError, zlib.error,
    ) as exc:
        raise TrackingValidationError("Il file XLSX non è valido.") from exc


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    _guard_xlsx(content)
    from openpyxl import load_workbook

    workbook = None
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False,
        )
        # pandas.read_excel(..., sheet_name=0) in the original always reads the
        # first worksheet, even when Excel saved another tab as active.
        sheet = workbook.worksheets[0]
        if sheet.max_column is not None and sheet.max_column > MAX_COLUMNS:
            raise TrackingLimitError(f"Il file supera il limite di {MAX_COLUMNS} colonne.")
        if sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1:
            raise TrackingLimitError(f"Il file supera il limite di {MAX_ROWS} righe.")
        iterator = sheet.iter_rows(values_only=True)
        columns = _unique_columns(next(iterator))
        rows = _bounded_rows(columns, iterator)
        return columns, rows
    except StopIteration as exc:
        raise TrackingValidationError("Il foglio XLSX è vuoto.") from exc
    except TrackingValidationError:
        raise
    except Exception as exc:
        raise TrackingValidationError("Il file XLSX non è leggibile.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _parse_xls(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    import xlrd

    workbook = None
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        sheet = workbook.sheet_by_index(0)
        if sheet.ncols > MAX_COLUMNS:
            raise TrackingLimitError(f"Il file supera il limite di {MAX_COLUMNS} colonne.")
        if sheet.nrows > MAX_ROWS + 1:
            raise TrackingLimitError(f"Il file supera il limite di {MAX_ROWS} righe.")
        if sheet.nrows * sheet.ncols > MAX_CELLS:
            raise TrackingLimitError(f"Il file supera il limite di {MAX_CELLS} celle.")
        columns = _unique_columns(sheet.row_values(0))
        rows = _bounded_rows(columns, (sheet.row_values(index) for index in range(1, sheet.nrows)))
        return columns, rows
    except TrackingValidationError:
        raise
    except Exception as exc:
        raise TrackingValidationError("Il file XLS non è leggibile.") from exc
    finally:
        if workbook is not None:
            workbook.release_resources()


def parse_tracking_file(file_name: str, content: bytes) -> dict:
    suffix = validate_file(file_name, content)
    columns, rows = (
        _parse_csv(content) if suffix == "csv"
        else _parse_xlsx(content) if suffix == "xlsx"
        else _parse_xls(content)
    )
    return {
        "file_name": Path(str(file_name).replace("\\", "/")).name,
        "row_count": len(rows),
        "columns": columns,
        "rows": rows,
        "detected": detect_tracking_columns(columns),
    }


def validate_mapping(mapping: dict[str, str], columns: list[str]) -> dict[str, str]:
    if set(mapping) - set(MAPPING_FIELDS):
        raise TrackingValidationError("La mappatura contiene campi non supportati.")
    result = {field: clean(mapping.get(field)) for field in MAPPING_FIELDS}
    existing = set(columns)
    if any(value and value not in existing for value in result.values()):
        raise TrackingValidationError("La mappatura fa riferimento a una colonna inesistente.")
    if not result["id_order_unit"] and not result["id_order"]:
        raise TrackingValidationError(
            "Associa almeno la colonna Numero ordine oppure ID unità ordine."
        )
    shipment_fields = ("carrier_code", "tracking_numbers", "combined_shipment")
    if not any(result[field] for field in shipment_fields):
        raise TrackingValidationError(
            "Associa almeno Tracking, Corriere oppure il campo combinato."
        )
    return result


def mapped_tracking_rows(rows: list[dict[str, str]], mapping: dict[str, str]):
    for index, record in enumerate(rows, start=1):
        unit_id = clean(record.get(mapping["id_order_unit"])) if mapping["id_order_unit"] else ""
        order_id = clean(record.get(mapping["id_order"])) if mapping["id_order"] else ""
        if len(unit_id) > MAX_IDENTIFIER_LENGTH or len(order_id) > MAX_IDENTIFIER_LENGTH:
            raise TrackingValidationError(
                f"Gli identificativi ordine non possono superare {MAX_IDENTIFIER_LENGTH} caratteri."
            )
        carrier = clean(record.get(mapping["carrier_code"])) if mapping["carrier_code"] else ""
        tracking = clean(record.get(mapping["tracking_numbers"])) \
            if mapping["tracking_numbers"] else ""
        combined = clean(record.get(mapping["combined_shipment"])) \
            if mapping["combined_shipment"] else ""
        if combined:
            combined_carrier, combined_tracking = split_shipment_text(combined)
            carrier = carrier or combined_carrier
            tracking = tracking or combined_tracking
        yield {
            "row": index, "order_unit_id": unit_id, "order_id": order_id,
            "carrier": carrier, "tracking": normalize_tracking(tracking),
        }
