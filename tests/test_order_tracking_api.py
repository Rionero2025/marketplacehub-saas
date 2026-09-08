import asyncio
import io
import json
import tempfile
import threading
import zipfile
from collections import deque
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime
from uuid import UUID, uuid4

import httpx
import pytest
import starlette.formparsers
import test_tenancy_api
from fastapi.testclient import TestClient
from marketplace_hub_api import orders as orders_api
from marketplace_hub_api import tracking_upload as tracking_upload_api
from marketplace_hub_api.tracking_upload import (
    MAX_MULTIPART_BODY_BYTES,
    TrackingMultipartLimitError,
    TrackingMultipartTimeoutError,
    read_tracking_multipart,
)
from marketplace_hub_core.auth.schema import auth_sessions
from marketplace_hub_core.orders import repository as orders_repository
from marketplace_hub_core.orders import service as orders_service
from marketplace_hub_core.orders import tracking as tracking_core
from marketplace_hub_core.orders.admission import (
    AsyncTrackingAdmission,
    MemoryTrackingAdmission,
    RedisTrackingAdmission,
    TrackingAdmissionBusyError,
    TrackingAdmissionCapacityError,
    TrackingAdmissionUnavailableError,
)
from marketplace_hub_core.orders.repository import (
    OrdersAmbiguousMatchError,
    OrdersImportBusyError,
    SqlOrdersRepository,
)
from marketplace_hub_core.orders.schema import order_lines, order_tracking_events
from marketplace_hub_core.orders.tracking import (
    TrackingLimitError,
    TrackingValidationError,
    detect_tracking_columns,
    parse_tracking_file,
    split_shipment_text,
)
from marketplace_hub_core.seller_settings.schema import seller_marketplace_accounts
from marketplace_hub_core.tenancy.schema import membership_permissions
from openpyxl import Workbook
from redis.exceptions import RedisError
from sqlalchemy import create_engine, func, select
from sqlalchemy.dialects.postgresql import dialect as postgresql_dialect
from sqlalchemy.exc import OperationalError
from starlette.requests import Request
from test_orders_api import account, owner, path, read, row, run, start
from test_orders_api import configured as orders_configured

workspace = test_tenancy_api.workspace


@pytest.fixture
def tracking_workspace(workspace):
    return orders_configured.__wrapped__(workspace)


def tracking_path(seller, suffix):
    return path(seller, f"/tracking/{suffix}")


def upload(client, seller, account_id, suffix, content, mapping=None, file_name="orders.csv",
           environment="live"):
    data = {} if mapping is None else {"mapping": json.dumps(mapping)}
    return client.post(
        tracking_path(seller, suffix),
        params={"account_id": str(account_id), "environment": environment},
        files={"file": (file_name, content, "application/octet-stream")},
        data=data,
    )


def test_original_column_detection_combined_split_and_xlsx_parser():
    columns = ["Numero ordine", "ID unità ordine", "Spedito con", "Tracciabilità"]
    assert detect_tracking_columns(columns) == {
        "id_order": "Numero ordine", "id_order_unit": "ID unità ordine",
        "carrier_code": "Spedito con", "tracking_numbers": "Tracciabilità",
    }
    assert split_shipment_text("DPD | 08448875901263") == ("DPD", "08448875901263")
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Spedizione"])
    workbook.active.append(["ORDER-1", "DPD | TRACK-42"])
    content = io.BytesIO()
    workbook.save(content)
    parsed = parse_tracking_file("portal.xlsx", content.getvalue())
    assert parsed["row_count"] == 1
    assert parsed["detected"] == {
        "id_order": "Numero ordine", "combined_shipment": "Spedizione",
    }
    multi_sheet = Workbook()
    multi_sheet.active.append(["Numero ordine", "Tracking"])
    multi_sheet.active.append(["FIRST-SHEET", "TRACK-FIRST"])
    multi_sheet.create_sheet("Active but ignored").append(["Wrong", "Sheet"])
    multi_sheet.active = 1
    multi_content = io.BytesIO()
    multi_sheet.save(multi_content)
    multi_parsed = parse_tracking_file("multi.xlsx", multi_content.getvalue())
    assert multi_parsed["rows"] == [{
        "Numero ordine": "FIRST-SHEET", "Tracking": "TRACK-FIRST",
    }]
    malicious = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source, zipfile.ZipFile(
        malicious, "w",
    ) as destination:
        for entry in source.infolist():
            value = source.read(entry)
            if entry.filename == "xl/workbook.xml":
                value = b'<!DOCTYPE x [<!ENTITY leak SYSTEM "file:///etc/passwd">]>' + value
            destination.writestr(entry, value)
    with pytest.raises(TrackingValidationError, match="XML"):
        parse_tracking_file("malicious.xlsx", malicious.getvalue())
    with pytest.raises(TrackingLimitError):
        parse_tracking_file(
            "wide.csv", (",".join(f"c{x}" for x in range(101)) + "\n").encode(),
        )
    assert parse_tracking_file("empty.csv", b"Ordine;Tracking\n")["row_count"] == 0
    with pytest.raises(TrackingValidationError):
        parse_tracking_file("nul.csv", b"Ordine;Tracking\nORDER-1;A\x00B\n")


def test_csv_shape_is_bounded_before_reader_allocation_and_preserves_quoted_newlines(
    monkeypatch,
):
    reader_calls = []
    monkeypatch.setattr(
        tracking_core.csv.Sniffer, "sniff",
        lambda *_args, **_kwargs: tracking_core.csv.excel,
    )

    def forbidden_reader(*_args, **_kwargs):
        reader_calls.append(True)
        raise AssertionError("csv.reader must not receive an oversized record")

    monkeypatch.setattr(tracking_core.csv, "reader", forbidden_reader)
    with pytest.raises(TrackingLimitError, match="100 colonne"):
        parse_tracking_file("wide.csv", b"header\n" + (b"," * 500_000) + b"\n")
    assert reader_calls == []

    monkeypatch.undo()
    parsed = parse_tracking_file(
        "quoted.csv",
        b'Ordine;Tracking;Note\nORDER-1;"TRACK;ONE";"prima linea\nseconda linea"\n',
    )
    assert parsed["rows"] == [{
        "Ordine": "ORDER-1",
        "Tracking": "TRACK;ONE",
        "Note": "prima linea\nseconda linea",
    }]


def test_xlsx_encrypted_or_corrupt_entry_is_a_validation_error(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    content = io.BytesIO()
    workbook.save(content)
    original_open = zipfile.ZipFile.open

    def encrypted_entry(archive, name, *args, **kwargs):
        entry_name = name.filename if isinstance(name, zipfile.ZipInfo) else str(name)
        if entry_name == "xl/workbook.xml":
            raise RuntimeError("File is encrypted, password required")
        return original_open(archive, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "open", encrypted_entry)
    with pytest.raises(TrackingValidationError, match="XLSX non è valido"):
        parse_tracking_file("encrypted.xlsx", content.getvalue())


def test_xlsx_rejects_nonstandard_lzma_compression():
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_LZMA) as destination:
        destination.writestr("xl/workbook.xml", b"<workbook />")
    with pytest.raises(TrackingValidationError, match="compressione o cifratura"):
        parse_tracking_file("lzma.xlsx", archive.getvalue())


def test_xlsx_without_declared_dimensions_uses_iterative_limits():
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Export")
    sheet.append(["Numero ordine", "Tracking"])
    sheet.append(["ORDER-WRITE-ONLY", "TRACK-WRITE-ONLY"])
    content = io.BytesIO()
    workbook.save(content)

    parsed = parse_tracking_file("write-only.xlsx", content.getvalue())

    assert parsed["rows"] == [{
        "Numero ordine": "ORDER-WRITE-ONLY",
        "Tracking": "TRACK-WRITE-ONLY",
    }]


def _xlsx_with_global_table(content: bytes, name: str, value: bytes) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED,
    ) as destination:
        replaced = False
        for entry in source.infolist():
            if entry.filename == name:
                destination.writestr(entry, value)
                replaced = True
            else:
                destination.writestr(entry, source.read(entry))
        if not replaced:
            destination.writestr(name, value)
    return output.getvalue()


def test_xlsx_rejects_unused_shared_string_explosion_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    content = io.BytesIO()
    workbook.save(content)
    shared_strings = (
        b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        + (b"<si/>" * (tracking_core.MAX_XLSX_SHARED_STRINGS + 1))
        + b"</sst>"
    )
    hostile = _xlsx_with_global_table(
        content.getvalue(), "xl/sharedStrings.xml", shared_strings,
    )
    assert len(hostile) < tracking_core.MAX_FILE_BYTES
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive an oversized shared string table")

    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingLimitError, match="stringhe XLSX supera"):
        parse_tracking_file("hostile.xlsx", hostile)
    assert load_calls == []


def test_xlsx_rejects_shared_string_rich_text_explosion_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    content = io.BytesIO()
    workbook.save(content)
    shared_strings = (
        b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si>'
        + (b"<r><t/></r>" * 100_000)
        + b"</si></sst>"
    )
    hostile = _xlsx_with_global_table(
        content.getvalue(), "xl/sharedStrings.xml", shared_strings,
    )
    assert len(hostile) < tracking_core.MAX_FILE_BYTES
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive excessive rich text runs")

    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingLimitError, match="stringa XLSX contiene troppi elementi"):
        parse_tracking_file("hostile.xlsx", hostile)
    assert load_calls == []


def test_xlsx_rejects_style_explosion_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    content = io.BytesIO()
    workbook.save(content)
    styles = (
        b'<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        + (b"<xf/>" * tracking_core.MAX_XLSX_STYLE_ELEMENTS)
        + b"</styleSheet>"
    )
    hostile = _xlsx_with_global_table(content.getvalue(), "xl/styles.xml", styles)
    assert len(hostile) < tracking_core.MAX_FILE_BYTES
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive an oversized stylesheet")

    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingLimitError, match="stile XLSX contiene troppi"):
        parse_tracking_file("hostile.xlsx", hostile)
    assert load_calls == []


def test_xlsx_rejects_defined_name_explosion_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    content = io.BytesIO()
    workbook.save(content)
    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source:
        workbook_xml = source.read("xl/workbook.xml")
    defined_names = (
        b"<definedNames>"
        + (
            b'<definedName name="safe_name">Sheet!$A$1</definedName>'
            * (tracking_core.MAX_XLSX_DEFINED_NAMES + 1)
        )
        + b"</definedNames>"
    )
    workbook_xml = workbook_xml.replace(
        b"</workbook>", defined_names + b"</workbook>",
    )
    hostile = _xlsx_with_global_table(
        content.getvalue(), "xl/workbook.xml", workbook_xml,
    )
    assert len(hostile) < tracking_core.MAX_FILE_BYTES
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive excessive defined names")

    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingLimitError, match="troppi nomi definiti"):
        parse_tracking_file("hostile.xlsx", hostile)
    assert load_calls == []


def test_xlsx_rejects_duplicate_cell_explosion_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    content = io.BytesIO()
    workbook.save(content)
    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source:
        worksheet_xml = source.read("xl/worksheets/sheet1.xml")
    duplicate_row = (
        b'<row r="2">'
        + (b'<c r="A2" t="inlineStr"><is><t>x</t></is></c>' * 100_000)
        + b"</row>"
    )
    worksheet_xml = worksheet_xml.replace(
        b"</sheetData>", duplicate_row + b"</sheetData>",
    )
    hostile = _xlsx_with_global_table(
        content.getvalue(), "xl/worksheets/sheet1.xml", worksheet_xml,
    )
    assert len(hostile) < tracking_core.MAX_FILE_BYTES
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive duplicate worksheet cells")

    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingValidationError, match="duplicate o non ordinate"):
        parse_tracking_file("hostile.xlsx", hostile)
    assert load_calls == []


def test_xlsx_guard_accepts_implicit_row_and_cell_coordinates():
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    content = io.BytesIO()
    workbook.save(content)
    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source:
        worksheet_xml = source.read("xl/worksheets/sheet1.xml")
    for coordinate in (b"1", b"2"):
        worksheet_xml = worksheet_xml.replace(
            b'<row r="' + coordinate + b'">', b"<row>",
        )
    for coordinate in (b"A1", b"B1", b"A2", b"B2"):
        worksheet_xml = worksheet_xml.replace(
            b'<c r="' + coordinate + b'"', b"<c",
        )
    implicit = _xlsx_with_global_table(
        content.getvalue(), "xl/worksheets/sheet1.xml", worksheet_xml,
    )

    parsed = parse_tracking_file("implicit.xlsx", implicit)

    assert parsed["rows"] == [{
        "Numero ordine": "ORDER-1", "Tracking": "TRACK-1",
    }]


def test_xlsx_only_applies_tracking_shape_to_first_worksheet():
    workbook = Workbook(write_only=True)
    first = workbook.create_sheet("Tracking")
    first.append(["Numero ordine", "Tracking"])
    first.append(["ORDER-1", "TRACK-1"])
    second = workbook.create_sheet("Archivio non importato")
    second.append(["Dato"])
    for index in range(tracking_core.MAX_ROWS + 1):
        second.append([index])
    content = io.BytesIO()
    workbook.save(content)

    parsed = parse_tracking_file("multi-sheet.xlsx", content.getvalue())

    assert parsed["rows"] == [{
        "Numero ordine": "ORDER-1", "Tracking": "TRACK-1",
    }]


def test_xlsx_rejects_unselected_sheet_preamble_explosion_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    workbook.create_sheet("Secondo foglio").append(["Dato"])
    content = io.BytesIO()
    workbook.save(content)
    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source:
        worksheet_xml = source.read("xl/worksheets/sheet2.xml")
    worksheet_xml = worksheet_xml.replace(
        b"<dimension", (b"<extLst><ext/></extLst>" * 200_000) + b"<dimension",
    )
    hostile = _xlsx_with_global_table(
        content.getvalue(), "xl/worksheets/sheet2.xml", worksheet_xml,
    )
    assert len(hostile) < tracking_core.MAX_FILE_BYTES
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not parse an oversized sheet preamble")

    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingLimitError, match="preambolo.*troppi elementi"):
        parse_tracking_file("hostile.xlsx", hostile)
    assert load_calls == []


def test_xlsx_rejects_huge_xml_attribute_before_sax_and_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    content = io.BytesIO()
    workbook.save(content)
    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source:
        worksheet_xml = source.read("xl/worksheets/sheet1.xml")
    extension = b'<extLst><ext value="' + (b"x" * 1024 * 1024) + b'"/></extLst>'
    worksheet_xml = worksheet_xml.replace(
        b"<sheetData>", extension + b"<sheetData>",
    )
    hostile = _xlsx_with_global_table(
        content.getvalue(), "xl/worksheets/sheet1.xml", worksheet_xml,
    )
    assert len(hostile) < tracking_core.MAX_FILE_BYTES
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive an oversized XML attribute")

    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingLimitError, match="attributo XML troppo grande"):
        parse_tracking_file("hostile.xlsx", hostile)
    assert load_calls == []


def test_xlsx_lexical_guard_tracks_comment_terminator_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    content = io.BytesIO()
    workbook.save(content)
    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source:
        worksheet_xml = source.read("xl/worksheets/sheet1.xml")
    short_comment_xml = worksheet_xml.replace(
        b"<sheetData>", b"<!-- > is data, not the terminator --> <sheetData>",
    )
    short_comment = _xlsx_with_global_table(
        content.getvalue(), "xl/worksheets/sheet1.xml", short_comment_xml,
    )
    assert parse_tracking_file("comment.xlsx", short_comment)["row_count"] == 1

    long_comment_xml = worksheet_xml.replace(
        b"<sheetData>", b"<!-- >" + (b"x" * 1024 * 1024) + b" --> <sheetData>",
    )
    hostile = _xlsx_with_global_table(
        content.getvalue(), "xl/worksheets/sheet1.xml", long_comment_xml,
    )
    assert len(hostile) < tracking_core.MAX_FILE_BYTES
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive an oversized XML comment")

    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingLimitError, match="elemento XML troppo grande"):
        parse_tracking_file("hostile.xlsx", hostile)
    assert load_calls == []


def test_xlsx_enforces_logical_cell_grid_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    content = io.BytesIO()
    workbook.save(content)
    monkeypatch.setattr(tracking_core, "MAX_CELLS", 4)
    assert parse_tracking_file("at-limit.xlsx", content.getvalue())["row_count"] == 1

    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive an oversized logical grid")

    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    monkeypatch.setattr(tracking_core, "MAX_CELLS", 3)
    with pytest.raises(TrackingLimitError, match="limite di 3 celle"):
        parse_tracking_file("over-limit.xlsx", content.getvalue())
    assert load_calls == []

    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source:
        worksheet_xml = source.read("xl/worksheets/sheet1.xml")
    worksheet_xml = worksheet_xml.replace(b'A1:B2', b'A1:CV10001')
    sparse = _xlsx_with_global_table(
        content.getvalue(), "xl/worksheets/sheet1.xml", worksheet_xml,
    )
    monkeypatch.setattr(tracking_core, "MAX_CELLS", 200_000)
    with pytest.raises(TrackingLimitError, match="limite di 200000 celle"):
        parse_tracking_file("sparse-over-limit.xlsx", sparse)
    assert load_calls == []


def test_xlsx_counts_text_outside_cells_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    content = io.BytesIO()
    workbook.save(content)
    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source:
        worksheet_xml = source.read("xl/worksheets/sheet1.xml")
    worksheet_xml = worksheet_xml.replace(b"<sheetData>", (b" " * 65) + b"<sheetData>")
    hostile = _xlsx_with_global_table(
        content.getvalue(), "xl/worksheets/sheet1.xml", worksheet_xml,
    )
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive excessive worksheet text")

    monkeypatch.setattr(tracking_core, "MAX_XLSX_WORKSHEET_TEXT", 64)
    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingLimitError, match="foglio XLSX contiene troppi dati"):
        parse_tracking_file("hostile.xlsx", hostile)
    assert load_calls == []


def test_csv_and_xls_enforce_logical_cell_cap(monkeypatch):
    monkeypatch.setattr(tracking_core, "MAX_CELLS", 4)
    assert parse_tracking_file(
        "at-limit.csv", b"Ordine;Tracking\nORDER-1;TRACK-1\n",
    )["row_count"] == 1
    with pytest.raises(TrackingLimitError, match="limite di 4 celle"):
        parse_tracking_file(
            "over-limit.csv",
            b"Ordine;Tracking\nORDER-1;TRACK-1\nORDER-2;TRACK-2\n",
        )

    class FakeSheet:
        ncols = 2
        nrows = 3

    class FakeBook:
        @staticmethod
        def sheet_by_index(_index):
            return FakeSheet()

        @staticmethod
        def release_resources():
            return None

    import xlrd

    monkeypatch.setattr(xlrd, "open_workbook", lambda **_kwargs: FakeBook())
    with pytest.raises(TrackingLimitError, match="limite di 4 celle"):
        parse_tracking_file("over-limit.xls", b"valid-biff-fixture-boundary")


def test_xlsx_enforces_cumulative_xml_element_budget_before_openpyxl(monkeypatch):
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Tracking"])
    workbook.active.append(["ORDER-1", "TRACK-1"])
    content = io.BytesIO()
    workbook.save(content)
    hostile = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content.getvalue())) as source, zipfile.ZipFile(
        hostile, "w", compression=zipfile.ZIP_DEFLATED,
    ) as destination:
        for entry in source.infolist():
            destination.writestr(entry, source.read(entry))
        small_part = b"<root>" + (b"<item/>" * 100) + b"</root>"
        for index in range(30):
            destination.writestr(f"customXml/item{index}.xml", small_part)
    load_calls = []

    def forbidden_load(*_args, **_kwargs):
        load_calls.append(True)
        raise AssertionError("openpyxl must not receive excessive aggregate XML")

    monkeypatch.setattr(tracking_core, "MAX_XLSX_XML_ELEMENTS_TOTAL", 2_500)
    monkeypatch.setattr("openpyxl.load_workbook", forbidden_load)
    with pytest.raises(TrackingLimitError, match="elementi XML complessivi"):
        parse_tracking_file("hostile.xlsx", hostile.getvalue())
    assert load_calls == []


def test_capabilities_are_readable_but_writes_require_logistics_modify(tracking_workspace):
    client, seller, organization, _, _ = owner(tracking_workspace, read_only=True)
    account_id = account(tracking_workspace, seller, organization)
    response = client.get(
        tracking_path(seller, "capabilities"), params={"account_id": str(account_id)},
    )
    assert response.status_code == 200
    assert response.json()["capabilities"] == {
        "marketplace": "kaufland", "formats": ["csv", "xlsx", "xls"],
        "max_bytes": 5 * 1024 * 1024, "max_rows": 10_000,
        "max_columns": 100, "max_cells": 200_000,
        "max_cell_length": 2_000, "preview_rows": 10,
        "can_import": False, "can_edit": False,
    }
    denied = upload(client, seller, account_id, "preview", b"Numero ordine\nORDER-1\n")
    assert denied.status_code == 403


def test_preview_is_bounded_and_validates_format_mapping_and_marketplace(tracking_workspace):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    content = "Numero ordine;Spedizione\n" + "\n".join(
        f"ORDER-{number};DPD | TRACK-{number}" for number in range(12)
    )
    response = upload(client, seller, account_id, "preview", content.encode())
    assert response.status_code == 200
    assert response.json()["status"] == "ready_for_mapping"
    preview = response.json()["preview"]
    assert preview["row_count"] == 12 and len(preview["rows"]) == 10
    assert preview["detected"] == {
        "id_order": "Numero ordine", "combined_shipment": "Spedizione",
    }
    unsupported = upload(
        client, seller, account_id, "preview", b"x", file_name="orders.pdf",
    )
    assert unsupported.status_code == 422
    oversized = upload(
        client, seller, account_id, "preview", b"x" * (5 * 1024 * 1024 + 1),
    )
    assert oversized.status_code == 413
    invalid_mapping = upload(
        client, seller, account_id, "import", content.encode(),
        {"id_order": "missing", "combined_shipment": "Spedizione"},
    )
    assert invalid_mapping.status_code == 422
    worten_id = account(tracking_workspace, seller, organization, marketplace="worten")
    assert client.get(tracking_path(seller, "capabilities"),
                      params={"account_id": str(worten_id)}).status_code == 422


def test_portal_import_updates_all_units_unit_id_wins_and_writes_audit(tracking_workspace):
    client, seller, organization, _, user = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [
        row("unit-1", order_id="ORDER-1", details={"carrier": "", "tracking": ""}),
        row("unit-2", order_id="ORDER-1", details={"carrier": "", "tracking": ""}),
    ]
    run(tracking_workspace, start(client, seller, account_id))
    mapping = {"id_order": "Numero ordine", "combined_shipment": "Spedizione"}
    imported = upload(
        client, seller, account_id, "import",
        b"Numero ordine;Spedizione\nORDER-1;DPD | TRACK-42\nUNKNOWN;GLS | LOST\n;\n",
        mapping,
    )
    assert imported.status_code == 200
    assert imported.json()["status"] == "completed_with_warnings"
    assert imported.json()["result"] == {
        "updated": 2,
        "unmatched": [{"row": 2, "order_unit_id": "", "order_id": "UNKNOWN"}],
        "invalid": [{"row": 3, "error": "Identificativo ordine assente"}],
    }
    items = read(client, seller, account_id).json()["items"]
    assert all(item["details"]["tracking"] == "TRACK-42" for item in items)
    assert all(item["details"]["tracking_source"] == "portal_import" for item in items)
    with tracking_workspace.engine.connect() as connection:
        events = connection.execute(select(order_tracking_events)).mappings().all()
    assert len(events) == 2
    assert all(
        event["actor_id"] == user.id and event["source"] == "portal_import"
        for event in events
    )

    unit_mapping = {
        "id_order_unit": "Unità", "id_order": "Ordine", "carrier_code": "Corriere",
    }
    exact = upload(
        client, seller, account_id, "import",
        b"Unit\xe0;Ordine;Corriere\nunit-1;WRONG;GLS\n", unit_mapping,
    )
    assert exact.json()["result"]["updated"] == 1
    by_unit = {item["external_line_id"]: item for item in read(
        client, seller, account_id,
    ).json()["items"]}
    assert by_unit["unit-1"]["details"]["carrier"] == "GLS"
    assert by_unit["unit-2"]["details"]["carrier"] == "DPD"


def test_manual_patch_preserves_blank_field_and_survives_resync(tracking_workspace):
    client, seller, organization, _, user = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [
        row(details={"carrier": "DPD", "tracking": "API-TRACK"}),
    ]
    run(tracking_workspace, start(client, seller, account_id))
    item = read(client, seller, account_id).json()["items"][0]
    changed = client.patch(tracking_path(seller, "manual"), json={
        "account_id": str(account_id), "environment": "live", "line_id": item["id"],
        "carrier": "GLS", "tracking": "",
    })
    assert changed.status_code == 200 and changed.json()["updated"] == 1
    details = changed.json()["item"]["details"]
    assert details["carrier"] == "GLS" and details["tracking"] == "API-TRACK"
    assert details["carrier_source"] == "manual"
    assert details["carrier_updated_by"] == str(user.id)

    tracking_workspace.fetcher.items = [row(details={"carrier": "", "tracking": ""})]
    run(tracking_workspace, start(client, seller, account_id))
    saved = read(client, seller, account_id).json()["items"][0]["details"]
    assert saved["carrier"] == "GLS" and saved["carrier_source"] == "manual"
    assert saved["tracking"] == "API-TRACK" and saved["tracking_source"] == "api"
    assert client.patch(tracking_path(seller, "manual"), json={
        "account_id": str(account_id), "line_id": str(UUID(int=0)),
        "carrier": "DHL", "tracking": "",
    }).status_code == 404
    assert client.patch(tracking_path(seller, "manual"), json={
        "account_id": str(account_id), "line_id": item["id"],
        "carrier": "", "tracking": "", "unexpected": str(uuid4()),
    }).status_code == 422
    assert client.patch(tracking_path(seller, "manual"), json={
        "account_id": str(account_id), "line_id": item["id"],
        "carrier": "", "tracking": "",
    }).status_code == 422
    assert client.patch(tracking_path(seller, "manual"), json={
        "account_id": str(account_id), "line_id": item["id"],
        "carrier": "DHL\u0000", "tracking": "TRACK",
    }).status_code == 422


def test_tracking_scope_isolates_environment_and_seller(tracking_workspace):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [row(details={"carrier": "", "tracking": ""})]
    run(tracking_workspace, start(client, seller, account_id))
    run(tracking_workspace, start(client, seller, account_id, environment="playground"))
    upload(client, seller, account_id, "import", b"Ordine;Tracking\nORDER-1;LIVE-ONLY\n",
           {"id_order": "Ordine", "tracking_numbers": "Tracking"})
    assert read(client, seller, account_id).json()["items"][0]["details"]["tracking"] == "LIVE-ONLY"
    assert not read(client, seller, account_id,
                    environment="playground").json()["items"][0]["details"]["tracking"]

    foreign, foreign_seller, foreign_org, _, _ = owner(tracking_workspace)
    foreign_account = account(tracking_workspace, foreign_seller, foreign_org)
    assert foreign.get(tracking_path(seller, "capabilities"),
                       params={"account_id": str(account_id)}).status_code == 404
    assert client.get(tracking_path(seller, "capabilities"),
                      params={"account_id": str(foreign_account)}).status_code == 404


def test_authentication_precedes_multipart_and_json_parsing(tracking_workspace, monkeypatch):
    multipart_calls = []
    json_calls = []

    async def forbidden_multipart(*_args, **_kwargs):
        multipart_calls.append(True)
        raise AssertionError("multipart parsed before authentication")

    async def forbidden_json(*_args, **_kwargs):
        json_calls.append(True)
        raise AssertionError("json parsed before authentication")

    monkeypatch.setattr(orders_api, "read_tracking_multipart", forbidden_multipart)
    monkeypatch.setattr(orders_api, "read_tracking_json", forbidden_json)
    client = TestClient(tracking_workspace.app)
    seller_id, account_id = uuid4(), uuid4()
    preview = client.post(
        tracking_path(seller_id, "preview"), params={"account_id": str(account_id)},
        content=b"malformed", headers={"content-type": "multipart/form-data; boundary=x"},
    )
    manual = client.patch(
        tracking_path(seller_id, "manual"), content=b"malformed",
        headers={"content-type": "application/json"},
    )
    assert preview.status_code == manual.status_code == 401
    assert multipart_calls == [] and json_calls == []


def test_tracking_admission_bounds_actor_account_and_global_capacity():
    admission = MemoryTrackingAdmission(maximum=2)
    first = admission.acquire(uuid4(), uuid4(), uuid4(), "live")
    actor = uuid4()
    seller = uuid4()
    account_id = uuid4()
    second = admission.acquire(actor, seller, account_id, "live")
    with pytest.raises(TrackingAdmissionCapacityError):
        admission.acquire(uuid4(), uuid4(), uuid4(), "live")
    admission.release(first)
    with pytest.raises(TrackingAdmissionBusyError):
        admission.acquire(actor, uuid4(), uuid4(), "live")
    with pytest.raises(TrackingAdmissionBusyError):
        admission.acquire(uuid4(), seller, account_id, "playground")
    admission.release(second)
    replacement = admission.acquire(actor, seller, account_id, "live")
    admission.release(replacement)


def test_redis_tracking_admission_uses_atomic_cross_process_lease():
    class RedisStub:
        def __init__(self, outcomes):
            self.outcomes = deque(outcomes)
            self.calls = []

        def eval(self, *args):
            self.calls.append(args)
            if "ZREMRANGEBYSCORE" not in args[0]:
                return 1
            outcome = self.outcomes.popleft()
            if isinstance(outcome, Exception):
                raise outcome
            return outcome

    client = RedisStub([1, 0, -1])
    admission = RedisTrackingAdmission(client, maximum=2, ttl_ms=60_000)
    lease = admission.acquire(uuid4(), uuid4(), uuid4(), "live")
    acquire_call = client.calls[0]
    assert acquire_call[1] == 3
    assert all("{orders}" in key for key in acquire_call[2:5])
    assert acquire_call[-3:] == (lease.token, 2, 60_000)
    with pytest.raises(TrackingAdmissionBusyError):
        admission.acquire(uuid4(), uuid4(), uuid4(), "live")
    with pytest.raises(TrackingAdmissionCapacityError):
        admission.acquire(uuid4(), uuid4(), uuid4(), "live")
    admission.release(lease)
    assert "ZREMRANGEBYSCORE" not in client.calls[-1][0]

    unavailable = RedisTrackingAdmission(RedisStub([RedisError("private redis URL")]))
    with pytest.raises(TrackingAdmissionUnavailableError):
        unavailable.acquire(uuid4(), uuid4(), uuid4(), "live")


def test_cancelled_shared_admission_releases_late_redis_and_local_leases():
    class BlockingSharedAdmission:
        maximum = 2

        def __init__(self):
            self.backing = MemoryTrackingAdmission(maximum=self.maximum)
            self.entered = threading.Event()
            self.proceed = threading.Event()
            self.released = threading.Event()
            self.calls = 0

        def acquire(self, *scope):
            self.calls += 1
            if self.calls == 1:
                self.entered.set()
                assert self.proceed.wait(timeout=5)
            return self.backing.acquire(*scope)

        def release(self, lease):
            self.backing.release(lease)
            self.released.set()

    async def scenario():
        shared = BlockingSharedAdmission()
        admission = AsyncTrackingAdmission(shared)
        scope = (uuid4(), uuid4(), uuid4(), "live")
        pending = asyncio.create_task(admission.acquire(*scope))
        assert await asyncio.to_thread(shared.entered.wait, 5)
        pending.cancel()
        with pytest.raises(asyncio.CancelledError):
            await pending
        shared.proceed.set()
        assert await asyncio.to_thread(shared.released.wait, 5)
        for _ in range(100):
            if admission.pending_cleanups == 0:
                break
            await asyncio.sleep(0.001)
        assert admission.pending_cleanups == 0
        replacement = await admission.acquire(*scope)
        await admission.release(replacement)

    asyncio.run(scenario())


def test_cancelled_release_finishes_shared_cleanup_without_task_leak():
    class BlockingReleaseAdmission:
        maximum = 2

        def __init__(self):
            self.backing = MemoryTrackingAdmission(maximum=self.maximum)
            self.entered = threading.Event()
            self.proceed = threading.Event()
            self.release_calls = 0

        def acquire(self, *scope):
            return self.backing.acquire(*scope)

        def release(self, lease):
            self.release_calls += 1
            if self.release_calls == 1:
                self.entered.set()
                assert self.proceed.wait(timeout=5)
            self.backing.release(lease)

    async def scenario():
        shared = BlockingReleaseAdmission()
        admission = AsyncTrackingAdmission(shared)
        scope = (uuid4(), uuid4(), uuid4(), "live")
        lease = await admission.acquire(*scope)
        pending = asyncio.create_task(admission.release(lease))
        assert await asyncio.to_thread(shared.entered.wait, 5)
        pending.cancel()
        with pytest.raises(asyncio.CancelledError):
            await pending
        shared.proceed.set()
        for _ in range(5_000):
            if admission.pending_cleanups == 0:
                break
            await asyncio.sleep(0.001)
        assert admission.pending_cleanups == 0
        replacement = await admission.acquire(*scope)
        await admission.release(replacement)

    asyncio.run(scenario())


def test_second_tracking_upload_is_rejected_before_body_read_or_parser(
    tracking_workspace, monkeypatch,
):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    original_reader = orders_api.read_tracking_multipart
    original_parser = orders_service.parse_tracking_file
    reader_entered = threading.Event()
    release_reader = threading.Event()
    reader_calls = []
    parser_calls = []

    async def blocking_reader(request, require_mapping):
        reader_calls.append(require_mapping)
        reader_entered.set()
        await asyncio.to_thread(release_reader.wait)
        return await original_reader(request, require_mapping=require_mapping)

    def counted_parser(file_name, content):
        parser_calls.append(file_name)
        return original_parser(file_name, content)

    monkeypatch.setattr(orders_api, "read_tracking_multipart", blocking_reader)
    monkeypatch.setattr(orders_service, "parse_tracking_file", counted_parser)
    with ThreadPoolExecutor(max_workers=1) as executor:
        first = executor.submit(
            upload, client, seller, account_id, "preview",
            b"Ordine;Tracking\nORDER-1;TRACK-1\n",
        )
        assert reader_entered.wait(timeout=5)
        try:
            second = upload(
                client, seller, account_id, "preview",
                b"Ordine;Tracking\nORDER-2;MUST-NOT-PARSE\n",
            )
            assert second.status_code == 409
            assert second.json()["detail"] == (
                "Un’altra operazione tracking è già in corso per questo account. "
                "Attendi e riprova."
            )
            assert reader_calls == [False]
            assert parser_calls == []
        finally:
            release_reader.set()
        assert first.result(timeout=5).status_code == 200
    assert reader_calls == [False]
    assert parser_calls == ["orders.csv"]


def test_upload_timeout_releases_admission_for_the_next_request(
    tracking_workspace, monkeypatch,
):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    original_reader = orders_api.read_tracking_multipart
    calls = []

    async def timeout_once(request, require_mapping):
        calls.append(require_mapping)
        if len(calls) == 1:
            raise TrackingMultipartTimeoutError(
                "Tempo massimo di caricamento superato. Riprova."
            )
        return await original_reader(request, require_mapping=require_mapping)

    monkeypatch.setattr(orders_api, "read_tracking_multipart", timeout_once)
    timed_out = upload(
        client, seller, account_id, "preview",
        b"Ordine;Tracking\nORDER-1;TRACK-1\n",
    )
    assert timed_out.status_code == 408
    assert timed_out.json()["detail"] == "Tempo massimo di caricamento superato. Riprova."
    retried = upload(
        client, seller, account_id, "preview",
        b"Ordine;Tracking\nORDER-1;TRACK-1\n",
    )
    assert retried.status_code == 200
    assert calls == [False, False]


def test_multipart_reader_times_out_a_pending_stream(monkeypatch):
    monkeypatch.setattr(tracking_upload_api, "TRACKING_BODY_TIMEOUT_SECONDS", 0.01)

    async def receive():
        await asyncio.Event().wait()

    request = Request({
        "type": "http", "method": "POST", "path": "/", "query_string": b"",
        "headers": [(b"content-type", b"multipart/form-data; boundary=slow")],
    }, receive)
    with pytest.raises(TrackingMultipartTimeoutError, match="Tempo massimo"):
        asyncio.run(read_tracking_multipart(request, require_mapping=False))


def test_stream_reader_rejects_oversize_body_before_consuming_remainder():
    chunks = deque([b"\r\n" * (32 * 1024)] * 84 + [b"must-not-be-consumed"])
    initial_count = len(chunks)

    async def receive():
        chunk = chunks.popleft()
        return {"type": "http.request", "body": chunk, "more_body": bool(chunks)}

    request = Request({
        "type": "http", "method": "POST", "path": "/", "query_string": b"",
        "headers": [(b"content-type", b"multipart/form-data; boundary=x")],
    }, receive)
    with pytest.raises(TrackingMultipartLimitError):
        asyncio.run(read_tracking_multipart(request, require_mapping=False))
    assert len(chunks) > 0
    assert initial_count - len(chunks) <= (MAX_MULTIPART_BODY_BYTES // (64 * 1024)) + 1

    huge_length = Request({
        "type": "http", "method": "POST", "path": "/", "query_string": b"",
        "headers": [
            (b"content-type", b"multipart/form-data; boundary=x"),
            (b"content-length", b"9" * 5_000),
        ],
    })
    with pytest.raises(TrackingMultipartLimitError):
        asyncio.run(read_tracking_multipart(huge_length, require_mapping=False))


def test_multipart_rejects_extra_parts_and_manual_json_is_bounded(tracking_workspace):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    duplicate_files = client.post(
        tracking_path(seller, "preview"), params={"account_id": str(account_id)},
        files=[
            ("file", ("one.csv", b"Ordine\nONE\n", "text/csv")),
            ("file", ("two.csv", b"Ordine\nTWO\n", "text/csv")),
        ],
    )
    unexpected_preview_field = client.post(
        tracking_path(seller, "preview"), params={"account_id": str(account_id)},
        files={"file": ("one.csv", b"Ordine\nONE\n", "text/csv")},
        data={"unexpected": "value"},
    )
    unexpected_import_field = client.post(
        tracking_path(seller, "import"), params={"account_id": str(account_id)},
        files={"file": ("one.csv", b"Ordine;Tracking\nONE;T\n", "text/csv")},
        data={"mapping": "{}", "unexpected": "value"},
    )
    assert duplicate_files.status_code == 422
    assert unexpected_preview_field.status_code == 422
    assert unexpected_import_field.status_code == 422
    long_name = upload(
        client, seller, account_id, "preview", b"Ordine\nONE\n",
        file_name=("x" * 252) + ".csv",
    )
    path_name = upload(
        client, seller, account_id, "preview", b"Ordine\nONE\n",
        file_name="folder/orders.csv",
    )
    assert long_name.status_code == path_name.status_code == 422
    truncated = client.post(
        tracking_path(seller, "preview"), params={"account_id": str(account_id)},
        content=(
            b"--cut\r\nContent-Disposition: form-data; name=\"file\"; "
            b"filename=\"orders.csv\"\r\n\r\nOrdine\nORDER-1\n"
        ),
        headers={"content-type": "multipart/form-data; boundary=cut"},
    )
    assert truncated.status_code == 422
    oversized_manual = client.patch(
        tracking_path(seller, "manual"), content=b"x" * (8 * 1024 + 1),
        headers={"content-type": "application/json"},
    )
    assert oversized_manual.status_code == 413


def test_csv_xlsx_and_xls_stay_in_memory_at_http_boundary(
    tracking_workspace, monkeypatch,
):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    workbook = Workbook()
    workbook.active.append(["Numero ordine", "Spedizione"])
    workbook.active.append(["ORDER-1", "DPD | TRACK-1"])
    xlsx = io.BytesIO()
    workbook.save(xlsx)

    class FakeSheet:
        ncols = 2
        nrows = 2

        @staticmethod
        def row_values(index):
            return (["Numero ordine", "Spedizione"] if index == 0
                    else ["ORDER-1", "DPD | TRACK-1"])

    class FakeBook:
        @staticmethod
        def sheet_by_index(_index):
            return FakeSheet()

        @staticmethod
        def release_resources():
            return None

    import xlrd

    monkeypatch.setattr(xlrd, "open_workbook", lambda **_kwargs: FakeBook())

    def forbidden_tempfile(*_args, **_kwargs):
        raise AssertionError("tracking upload attempted to spool to disk")

    monkeypatch.setattr(starlette.formparsers, "SpooledTemporaryFile", forbidden_tempfile)
    monkeypatch.setattr(tempfile, "SpooledTemporaryFile", forbidden_tempfile)
    monkeypatch.setattr(tempfile, "TemporaryFile", forbidden_tempfile)
    uploads = (
        ("orders.csv", b"Numero ordine;Spedizione\nORDER-1;DPD | TRACK-1\n"),
        ("orders.xlsx", xlsx.getvalue()),
        ("orders.xls", b"valid-biff-fixture-boundary"),
    )
    for file_name, content in uploads:
        response = upload(
            client, seller, account_id, "preview", content, file_name=file_name,
        )
        assert response.status_code == 200
        assert response.json()["preview"]["row_count"] == 1


def test_import_runs_off_event_loop_thread(tracking_workspace, monkeypatch):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    token = client.cookies.get(tracking_workspace.settings.session_cookie_name)
    event_loop_thread = threading.get_ident()
    service_threads = []
    started = threading.Event()
    release = threading.Event()
    timed_out = []

    def imported(*_args, **_kwargs):
        service_threads.append(threading.get_ident())
        started.set()
        timed_out.append(not release.wait(2))
        return {"status": "completed", "result": {
            "updated": 0, "unmatched": [], "invalid": [],
        }}

    monkeypatch.setattr(tracking_workspace.orders, "import_tracking", imported)

    async def request_import():
        transport = httpx.ASGITransport(app=tracking_workspace.app)
        async with httpx.AsyncClient(
            transport=transport, base_url="http://testserver",
            cookies={tracking_workspace.settings.session_cookie_name: token},
        ) as async_client:
            import_request = asyncio.create_task(async_client.post(
                tracking_path(seller, "import"), params={"account_id": str(account_id)},
                files={"file": ("orders.csv", b"Ordine;Tracking\nORDER-1;T\n")},
                data={"mapping": json.dumps({
                    "id_order": "Ordine", "tracking_numbers": "Tracking",
                })},
            ))
            assert await asyncio.to_thread(started.wait, 1)
            try:
                readiness = await async_client.get("/health/ready")
            finally:
                release.set()
            return await import_request, readiness

    response, readiness = asyncio.run(request_import())
    assert response.status_code == 200
    assert readiness.status_code == 200
    assert service_threads and service_threads[0] != event_loop_thread
    assert timed_out == [False]


def test_repeated_targets_are_last_row_wins_and_retry_has_no_duplicate_audit(
    tracking_workspace,
):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [
        row("unit-1", details={"carrier": "", "tracking": ""}),
        row("unit-2", details={"carrier": "", "tracking": ""}),
    ]
    run(tracking_workspace, start(client, seller, account_id))
    content = (
        b"Ordine;Corriere;Tracking\n"
        b"ORDER-1;DPD;FIRST\n"
        b"ORDER-1;GLS;LAST\n"
    )
    mapping = {
        "id_order": "Ordine", "carrier_code": "Corriere",
        "tracking_numbers": "Tracking",
    }
    first = upload(client, seller, account_id, "import", content, mapping)
    assert first.json()["result"]["updated"] == 2
    items = read(client, seller, account_id).json()["items"]
    assert all(item["details"]["carrier"] == "GLS" for item in items)
    assert all(item["details"]["tracking"] == "LAST" for item in items)
    with tracking_workspace.engine.connect() as connection:
        first_events = connection.scalar(select(func.count()).select_from(order_tracking_events))
    assert first_events == 2
    retry = upload(client, seller, account_id, "import", content, mapping)
    assert retry.json()["result"]["updated"] == 0
    with tracking_workspace.engine.connect() as connection:
        retried_events = connection.scalar(select(func.count()).select_from(order_tracking_events))
    assert retried_events == first_events

    unit_rows = b"Unita;Corriere\nunit-1;UPS\nunit-1;DHL\n"
    unit_result = upload(
        client, seller, account_id, "import", unit_rows,
        {"id_order_unit": "Unita", "carrier_code": "Corriere"},
    )
    assert unit_result.json()["result"]["updated"] == 1
    by_unit = {item["external_line_id"]: item for item in read(
        client, seller, account_id,
    ).json()["items"]}
    assert by_unit["unit-1"]["details"]["carrier"] == "DHL"


def test_oversized_identifier_header_and_target_set_write_nothing(
    tracking_workspace, monkeypatch,
):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [
        row("unit-1", details={"carrier": "", "tracking": ""}),
        row("unit-2", details={"carrier": "", "tracking": ""}),
    ]
    run(tracking_workspace, start(client, seller, account_id))
    mapping = {"id_order": "Ordine", "tracking_numbers": "Tracking"}
    long_identifier = ("X" * 201).encode()
    mixed = upload(
        client, seller, account_id, "import",
        b"Ordine;Tracking\nORDER-1;WOULD-WRITE\n" + long_identifier + b";UNKNOWN\n",
        mapping,
    )
    assert mixed.status_code == 422
    long_header = ("H" * 201).encode()
    oversized_header = upload(
        client, seller, account_id, "import",
        b"Ordine;Tracking;" + long_header + b"\nORDER-1;WOULD-WRITE;x\n", mapping,
    )
    assert oversized_header.status_code == 413
    monkeypatch.setattr(orders_repository, "MAX_TRACKING_TARGET_UNITS", 1)
    amplified = upload(
        client, seller, account_id, "import",
        b"Ordine;Tracking\nORDER-1;WOULD-WRITE\n", mapping,
    )
    assert amplified.status_code == 413
    saved = read(client, seller, account_id).json()["items"]
    assert all(not item["details"].get("tracking") for item in saved)
    with tracking_workspace.engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(order_tracking_events)) == 0


@pytest.mark.parametrize("revocation", ["session", "permission", "account"])
def test_import_reauthorizes_after_parsing_before_any_write(
    tracking_workspace, monkeypatch, revocation,
):
    client, seller, organization, membership, user = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [
        row("unit-1", details={"carrier": "", "tracking": ""}),
    ]
    run(tracking_workspace, start(client, seller, account_id))
    original_parse = orders_service.parse_tracking_file

    def parse_then_revoke(file_name, content):
        parsed = original_parse(file_name, content)
        with tracking_workspace.engine.begin() as connection:
            if revocation == "session":
                connection.execute(auth_sessions.update().where(
                    auth_sessions.c.user_id == user.id,
                ).values(revoked_at=datetime.now(UTC)))
            elif revocation == "permission":
                connection.execute(membership_permissions.delete().where(
                    membership_permissions.c.membership_id == membership,
                    membership_permissions.c.permission_code == "LOGISTICS",
                ))
            else:
                connection.execute(seller_marketplace_accounts.update().where(
                    seller_marketplace_accounts.c.id == account_id,
                ).values(active=False))
        return parsed

    monkeypatch.setattr(orders_service, "parse_tracking_file", parse_then_revoke)
    response = upload(
        client, seller, account_id, "import",
        b"Ordine;Tracking\nORDER-1;MUST-NOT-WRITE\n",
        {"id_order": "Ordine", "tracking_numbers": "Tracking"},
    )
    assert response.status_code == {"session": 401, "permission": 403, "account": 422}[
        revocation
    ]
    with tracking_workspace.engine.connect() as connection:
        canonical = json.loads(connection.scalar(select(
            order_lines.c.canonical_json,
        ).where(order_lines.c.account_id == account_id)))
        assert not canonical["details"].get("tracking")
        assert connection.scalar(select(func.count()).select_from(order_tracking_events)) == 0


def test_manual_update_reauthenticates_after_bounded_json_read(tracking_workspace, monkeypatch):
    client, seller, organization, _, user = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [
        row("unit-1", details={"carrier": "", "tracking": ""}),
    ]
    run(tracking_workspace, start(client, seller, account_id))
    line_id = read(client, seller, account_id).json()["items"][0]["id"]
    original_reader = orders_api.read_tracking_json

    async def read_then_revoke(request):
        content = await original_reader(request)
        with tracking_workspace.engine.begin() as connection:
            connection.execute(auth_sessions.update().where(
                auth_sessions.c.user_id == user.id,
            ).values(revoked_at=datetime.now(UTC)))
        return content

    monkeypatch.setattr(orders_api, "read_tracking_json", read_then_revoke)
    response = client.patch(tracking_path(seller, "manual"), json={
        "account_id": str(account_id), "environment": "live", "line_id": line_id,
        "carrier": "GLS", "tracking": "MUST-NOT-WRITE",
    })
    assert response.status_code == 401
    with tracking_workspace.engine.connect() as connection:
        canonical = json.loads(connection.scalar(select(
            order_lines.c.canonical_json,
        ).where(order_lines.c.account_id == account_id)))
        assert not canonical["details"].get("tracking")
        assert connection.scalar(select(func.count()).select_from(order_tracking_events)) == 0


def test_import_rechecks_locked_account_before_writing(tracking_workspace, monkeypatch):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [
        row("unit-1", details={"carrier": "", "tracking": ""}),
    ]
    run(tracking_workspace, start(client, seller, account_id))
    original_lock = tracking_workspace.order_repository._lock_tracking_import

    def deactivate_then_lock(connection, organization_id, seller_id, locked_account_id):
        connection.execute(seller_marketplace_accounts.update().where(
            seller_marketplace_accounts.c.id == locked_account_id,
        ).values(active=False))
        return original_lock(connection, organization_id, seller_id, locked_account_id)

    monkeypatch.setattr(
        tracking_workspace.order_repository, "_lock_tracking_import", deactivate_then_lock,
    )
    response = upload(
        client, seller, account_id, "import",
        b"Ordine;Tracking\nORDER-1;MUST-NOT-WRITE\n",
        {"id_order": "Ordine", "tracking_numbers": "Tracking"},
    )
    assert response.status_code == 422
    assert response.json()["detail"] == "Verifica il marketplace prima di sincronizzare."
    with tracking_workspace.engine.connect() as connection:
        canonical = json.loads(connection.scalar(select(
            order_lines.c.canonical_json,
        ).where(order_lines.c.account_id == account_id)))
        assert not canonical["details"].get("tracking")
        assert connection.scalar(select(func.count()).select_from(order_tracking_events)) == 0


def test_ambiguous_legacy_unit_is_rejected_without_partial_update(tracking_workspace):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [
        row("shared-unit", order_id="ORDER-1", details={"carrier": "", "tracking": ""}),
    ]
    run(tracking_workspace, start(client, seller, account_id))
    with tracking_workspace.engine.begin() as connection:
        saved = dict(connection.execute(select(order_lines).where(
            order_lines.c.account_id == account_id,
        )).mappings().one())
        canonical = json.loads(saved["canonical_json"])
        canonical["order_id"] = "ORDER-2"
        saved.update(
            id=uuid4(), order_id="ORDER-2",
            canonical_json=json.dumps(canonical, ensure_ascii=False),
        )
        connection.execute(order_lines.insert().values(**saved))
    response = upload(
        client, seller, account_id, "import",
        b"Unita;Ordine;Corriere\n;ORDER-1;WOULD-WRITE\nshared-unit;;GLS\n",
        {
            "id_order_unit": "Unita", "id_order": "Ordine",
            "carrier_code": "Corriere",
        },
    )
    assert response.status_code == 422
    assert "più righe" in response.json()["detail"]
    assert all(
        not item["details"].get("carrier")
        for item in read(client, seller, account_id).json()["items"]
    )
    with tracking_workspace.engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(order_tracking_events)) == 0


def test_kaufland_sync_rejects_duplicate_unit_ids_atomically_under_account_lock(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'tracking-race.sqlite3'}")
    order_lines.create(engine)
    seller_marketplace_accounts.create(engine)
    repository = SqlOrdersRepository(engine)
    scope = {
        "organization_id": uuid4(), "seller_id": uuid4(), "account_id": uuid4(),
        "environment": "live", "marketplace": "kaufland",
    }
    now = datetime.now(UTC)
    with engine.begin() as connection:
        connection.execute(seller_marketplace_accounts.insert().values(
            id=scope["account_id"], seller_id=scope["seller_id"],
            organization_id=scope["organization_id"], marketplace="kaufland",
            account_name="sync-race", credentials_encrypted="encrypted",
            settings_json=json.dumps({
                "marketplace_hub_connection_v1": {"connection_status": "connected"},
            }),
            active=True, legacy_account_id=None, created_at=now, updated_at=now,
        ))

    with pytest.raises(OrdersAmbiguousMatchError):
        repository.upsert_batch(scope, [
            row("batch-duplicate", order_id="ORDER-BATCH-A"),
            row("batch-duplicate", order_id="ORDER-BATCH-B"),
        ])
    with engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(order_lines)) == 0

    barrier = threading.Barrier(2)

    def write(order_id):
        barrier.wait()
        try:
            repository.upsert_batch(scope, [row("shared-unit", order_id=order_id)])
            return "saved"
        except (OrdersAmbiguousMatchError, OrdersImportBusyError) as exc:
            return type(exc).__name__

    with ThreadPoolExecutor(max_workers=2) as executor:
        outcomes = list(executor.map(write, ("ORDER-A", "ORDER-B")))
    assert outcomes.count("saved") == 1
    with engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(order_lines)) == 1

    losing_order = "ORDER-A" if outcomes[0] != "saved" else "ORDER-B"
    with pytest.raises(OrdersAmbiguousMatchError):
        repository.upsert_batch(scope, [row("shared-unit", order_id=losing_order)])

    worten_scope = {**scope, "marketplace": "worten"}
    repository.upsert_batch(worten_scope, [
        row("shared-unit", order_id="WORTEN-A", marketplace="worten"),
        row("shared-unit", order_id="WORTEN-B", marketplace="worten"),
    ])
    with engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(order_lines)) == 3
    engine.dispose()


def test_tracking_import_scope_lock_fails_fast_and_commits_nothing_when_busy(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'tracking-lock.sqlite3'}")
    order_lines.create(engine)
    order_tracking_events.create(engine)
    seller_marketplace_accounts.create(engine)
    repository = SqlOrdersRepository(engine)
    now = datetime.now(UTC)
    scope = {
        "organization_id": uuid4(), "seller_id": uuid4(), "account_id": uuid4(),
        "environment": "live", "marketplace": "kaufland",
    }
    with engine.begin() as connection:
        connection.execute(seller_marketplace_accounts.insert().values(
            id=scope["account_id"], seller_id=scope["seller_id"],
            organization_id=scope["organization_id"], marketplace="kaufland",
            account_name="lock-test", credentials_encrypted="encrypted",
            settings_json=json.dumps({
                "marketplace_hub_connection_v1": {"connection_status": "connected"},
            }),
            active=True, legacy_account_id=None,
            created_at=now, updated_at=now,
        ))
    repository.upsert_batch(scope, [row("unit-lock", details={
        "carrier": "", "tracking": "",
    })])
    records = [{
        "row": 1, "order_unit_id": "unit-lock", "order_id": "",
        "carrier": "GLS", "tracking": "TRACK-LOCK",
    }]

    with engine.begin() as held:
        held.execute(seller_marketplace_accounts.update().where(
            seller_marketplace_accounts.c.id == scope["account_id"],
        ).values(updated_at=seller_marketplace_accounts.c.updated_at))
        with ThreadPoolExecutor(max_workers=1) as executor:
            future = executor.submit(
                repository.import_tracking_batch,
                scope["organization_id"], scope["seller_id"], scope["account_id"], "live",
                records=records, source="portal_import", actor_id=uuid4(),
            )
            with pytest.raises(OrdersImportBusyError):
                future.result(timeout=2)
        with engine.connect() as observer:
            assert observer.scalar(
                select(func.count()).select_from(order_tracking_events)
            ) == 0

    completed = repository.import_tracking_batch(
        scope["organization_id"], scope["seller_id"], scope["account_id"], "live",
        records=records, source="portal_import", actor_id=uuid4(),
    )
    assert completed["updated"] == 1
    engine.dispose()


def test_postgres_tracking_target_lock_is_nowait_and_maps_lock_conflict():
    captured = []

    class Locked(Exception):
        sqlstate = "55P03"

    class PostgreSQLConnection:
        dialect = type("Dialect", (), {"name": "postgresql"})()

        @staticmethod
        def execute(statement):
            captured.append(statement)
            raise OperationalError("SELECT", {}, Locked())

    with pytest.raises(OrdersImportBusyError):
        SqlOrdersRepository._execute_tracking_locked(
            PostgreSQLConnection(), select(order_lines),
        )
    sql = str(captured[0].compile(dialect=postgresql_dialect()))
    assert sql.endswith("FOR UPDATE NOWAIT")


def test_manual_and_batch_target_lock_conflicts_rollback_without_audit(
    tracking_workspace, monkeypatch,
):
    client, seller, organization, _, _ = owner(tracking_workspace)
    account_id = account(tracking_workspace, seller, organization)
    tracking_workspace.fetcher.items = [
        row("unit-locked", details={"carrier": "", "tracking": ""}),
    ]
    run(tracking_workspace, start(client, seller, account_id))
    line_id = read(client, seller, account_id).json()["items"][0]["id"]
    message = (
        "Un’altra importazione tracking è già in corso per questo account. "
        "Attendi e riprova."
    )

    def locked(*_args, **_kwargs):
        raise OrdersImportBusyError(message)

    monkeypatch.setattr(
        tracking_workspace.order_repository, "_execute_tracking_locked", locked,
    )
    manual = client.patch(tracking_path(seller, "manual"), json={
        "account_id": str(account_id), "environment": "live", "line_id": line_id,
        "carrier": "GLS", "tracking": "MUST-NOT-WRITE",
    })
    imported = upload(
        client, seller, account_id, "import",
        b"Ordine;Tracking\nORDER-1;MUST-NOT-WRITE\n",
        {"id_order": "Ordine", "tracking_numbers": "Tracking"},
    )
    assert manual.status_code == imported.status_code == 409
    assert manual.json()["detail"] == imported.json()["detail"] == message
    with tracking_workspace.engine.connect() as connection:
        canonical = json.loads(connection.scalar(select(
            order_lines.c.canonical_json,
        ).where(order_lines.c.account_id == account_id)))
        assert not canonical["details"].get("tracking")
        assert connection.scalar(select(func.count()).select_from(order_tracking_events)) == 0
